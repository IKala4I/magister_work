#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Звірка зібраного .docx із джерелом.

    python3 verify.py --src ../text/full.md --docx $DOCS/build/hourwell.docx --brief ../formatter-brief.md \\
                      --figures $DOCS/build/figures --private $DOCS/private/titulka.json

Що гарантує і чого НЕ гарантує. Перевірка на дослівність доводить, що дорогою від
full.md до .docx нічого не переписано. Вона нічого не каже про те, чи правильний
сам full.md: якщо розмітка зіпсувала джерело, обидві сторони міститимуть однакову
помилку й ідеально збігатимуться. Від цього захищають інші перевірки —
check_legend_symbols.py, audit_markdown.py і структурні перевірки таблиць у build.py.

Звірка йде від СИРОГО джерела — після погоджених правок реєстру, але до нормалізації
складальника. Кожна відмінність документа від джерела має пояснитися дозволеною заміною,
і дозволеність verify.py оцінює сам, власними правилами (judge), а не правилом build.py:
перевірка, що звіряє після того самого правила, не бачить його помилок. Так конвеєр,
відновлений 2026-09-28, друкував номер підрозділу «3.7» як «3,7», а звірка казала OK.

Перевірки:
  1. кожен непорожній рядок прози є в тексті документа (дозволено лише ' → ’, і не у
     вставках коду); код порівнюється без зняття розмітки й без жодних замін;
  2. кожна таблиця джерела — комірка до комірки проти таблиці документа; комірка або
     збігається, або відрізняється лише дозволеною заміною (кома й мінус — тільки в
     комірці, що вся є числом, і ніколи в номері підрозділу цього документа);
  3. кількість формул у джерелі = кількість об'єктів OMML (з урахуванням розбитих)
     і жодне кириличне слово з рядка формули не загубилося в LaTeX;
  4. рядки keep-verbatim із брифа — є в джерелі й у документі (апостроф і тире
     згортаються, як велить §7 брифа; структура рівняння — окремим винятком);
  5. кожна вставка з реєстру, яку застосовано, є в документі;
  6. кожен блок коду — рядок до рядка, разом із порожніми, проти абзаців коду документа
     (до 2026-09-28 складальник викидав порожні рядки з коду, а перевірка 1 їх не бачила);
  7. з --figures: кожен відрендерений рисунок стоїть на своєму місці — зображення перед підписом
     «Рис. N …» і кожен наступний аркуш після нього з підписом «Рис. N, аркуш k» —
     і це саме той рендер: SHA-256 байтів зображення в .docx збігається з файлом рендера;
     плейсхолдера такого рисунка в документі немає; зайвих зображень немає;
  8. з --private: кожен заповнений приватний плейсхолдер (титулка, анотації) у документі
     замінено значенням, незаповнений — лишився як є, і жодне приватне значення не потрапило
     у full.md (публічний репозиторій).
  9. оформлення за методичними рекомендаціями кафедри (2026), виміряне в готовому .docx, а не
     прочитане з намірів build.py: аркуш і поля, шрифт і інтервал, заголовок розділу у два рядки,
     відстань від заголовка до тексту, заголовок пункту в розрядці в підбір до тексту, підписи
     таблиць і рисунків, вільні рядки біля формул, заголовки додатків, перелік у дві колонки,
     типографіка (нерозривні пробіли перед тире й між числом і одиницею, нерозривні дефіси).
     `--self-test` псує зібраний документ по разу на кожне правило й вимагає, щоб перевірка
     це помітила.
Код виходу 1, якщо порушено будь-що з 1–9.
"""
import argparse, json, os, re, sys, zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build                                          # noqa: E402
from formulas import FORMULA_TEX, FORMULA_SPLIT, DOCX_EXEMPT   # noqa: E402
from docx import Document                              # noqa: E402

ws = lambda s: re.sub(r'\s+', ' ', s).strip()
fold = lambda s: s.replace("'", '\u2019').replace('—', '–')


def ref_cells(line):
    """НЕЗАЛЕЖНИЙ від build.py розбір комірок: лише правило GFM (скісна екранує
    наступний символ) і нічого більше. Навмисно не імпортується зі складальника:
    перевірка, що ділить код із тим, що перевіряє, не бачить його помилок —
    саме так пробіл у «за змінного |A_m(x)|» зникав непоміченим."""
    inner = line.strip()
    if inner.startswith('|'): inner = inner[1:]
    if inner.endswith('|') and not inner.endswith('\\|'): inner = inner[:-1]
    cells, cur, i = [], '', 0
    while i < len(inner):
        if inner[i] == '\\' and i + 1 < len(inner):
            cur += inner[i:i+2]; i += 2; continue
        if inner[i] == '|': cells.append(cur); cur = ''
        else: cur += inner[i]
        i += 1
    cells.append(cur)
    return [c.strip().replace('\\|', '|') for c in cells]


def unmath(s):
    """`$…$` → те, що читається в .docx: без знаків долара, підкреслень, «^» і фігурних дужок
    (індекс — це положення прогону, а не символ). Власне правило, не код build.py."""
    return re.sub(r'\$([^$\n]+)\$', lambda m: re.sub(r'[_^{}]', '', m.group(1)), s)


def strip_md(s):
    s = unmath(s)
    s = re.sub(r'^\s*```[a-z]*\s*', '', s)
    s = re.sub(r'\s*```\s*$', '', s)
    s = re.sub(r'\*\*(.+?)\*\*', r'\1', s)
    s = re.sub(r'(?<![\w\\])_([^_\n]+)_(?![\w])',
               lambda m: m.group(1) if build._is_prose_italic(m.group(1)) else m.group(0), s)
    s = s.replace('`', '')
    s = re.sub(r'^#{1,3}\s+', '', s)
    return ws(s)


def docx_text(path):
    doc = Document(path)
    out = [p.text for p in doc.paragraphs]

    def cells(tbl):
        for row in tbl.rows:
            for c in row.cells:
                out.extend(p.text for p in c.paragraphs)
                for t in c.tables: cells(t)
    for t in doc.tables: cells(t)
    with zipfile.ZipFile(path) as z:
        xml = z.read('word/document.xml').decode('utf-8')
    return ws(' '.join(out)), len(re.findall(r'<m:oMath>', xml))


CONTINUED = re.compile(r'Продовження таблиці ((?:\d+|[А-ЯІЇЄҐ])\.\d+)')


def docx_tables(path):
    """Таблиці даних документа по порядку, як [[текст комірки]]. Формули — теж таблиці,
    але без стилю «Table Grid», тож сюди не потрапляють. Таблиця, поділена на частини написом
    «Продовження таблиці N», зшивається назад: частина після такого напису додає до попередньої
    таблиці свої рядки без головки. Чи головка частини та сама, що в першої, перевіряє перевірка 9."""
    doc, out, prev = Document(path), [], ''
    for x in _body(doc):
        if _is_p(x):
            if x.text.strip(): prev = ws(x.text)
            continue
        if x.style is None or x.style.name != 'Table Grid': continue
        rows = [[ws(c.text) for c in row.cells] for row in x.rows]
        if CONTINUED.fullmatch(prev) and out: out[-1].extend(rows[1:])
        else: out.append(rows)
        prev = ''
    return out


def raw(n):
    """Рядок джерела з погодженими override-ами реєстру, але БЕЗ нормалізації складальника."""
    s = build.LINES[n - 1]
    o = build.OVERRIDE.get(n)
    if o and o['replace'] not in s: s = s.replace(o['find'], o['replace'], 1)
    return s


def code_lines(lines):
    """Власний облік рядків коду: огорожі ``` і все між ними, плюс діапазони коду з реєстру."""
    out, inside = set(), False
    for n, s in enumerate(lines, 1):
        if inside:
            out.add(n)
            if s.rstrip().endswith('```'): inside = False
        elif s.lstrip().startswith('```'):
            out.add(n); inside = s.count('```') < 2
    for a, b in build.CODE_RANGES: out.update(range(a, b + 1))
    return out


def heading_numbers(lines):
    """Номери розділів і підрозділів цього документа («3.7», «6.6.2»), прочитані з самих
    заголовків. Знання, незалежне від складальника: число, що збігається з номером
    підрозділу, — посилання на підрозділ, а не десятковий дріб. Якщо колись справжній дріб
    збіжеться з номером підрозділу, звірка впаде — тоді дріб у джерелі пишуть із комою,
    як усі інші."""
    return {m.group(1) for s in lines for m in [re.match(r'^#{1,3}\s+(\d+(?:\.\d+)+)\.?\s', s)] if m}


PROTECTED = re.compile(r'`[^`]*`|\S*(?:://|www\.|@)\S*|\S+\.(?:md|py|json|sql)\b\S*')


def protected(raw_text):
    """Фрагменти з апострофом, які бриф велить лишити як є: вставки в зворотних лапках і
    адреси (посилання, пошта, імена файлів) — `f'y_{t.id}'`, https://…/it's. Без лапок."""
    return [m.strip('`') for m in PROTECTED.findall(raw_text) if "'" in m]


def curl(raw_text):
    """Сирий рядок, у якому ' → ’ скрізь, крім захищених фрагментів: так має виглядати
    проза в документі, якщо апостроф замінено правильно."""
    spans = PROTECTED.findall(raw_text)
    return ''.join(part.replace("'", '’') + (spans[i] if i < len(spans) else '')
                   for i, part in enumerate(PROTECTED.split(raw_text)))


NAME_FIELDS = tuple(k for k, _, is_name in build.PRIVATE_FIELDS if is_name)
MIN_PT = 8.0                    # найдрібніший текст рисунка на папері (власник, 2026-09-29)
ALLOWED = {('.', ','): 'кома', ('-', '−'): 'мінус', ("'", '’'): 'апостроф'}
PURE_NUMBER = re.compile(r'[+\-−]?(\d+(?:\.\d+)?)(?:\s*±\s*\d+(?:\.\d+)?)?')


def judge(raw_cell, src, doc, sections):
    """Чи пояснюється відмінність комірки документа від сирого джерела дозволеною заміною.
    Повертає (причина, []) якщо ні, або (None, [види замін]) якщо так. Правила записано тут
    заново, словами брифа: кома й мінус — лише в комірці, що вся є числом, і ніколи в номері
    підрозділу; апостроф — будь-де, крім вставок у зворотних лапках."""
    if src == doc: return None, []
    if len(src) != len(doc): return 'змінилася довжина', []
    kinds = set()
    for x, y in zip(src, doc):
        if x != y:
            if (x, y) not in ALLOWED: return f'заміна {x!r} → {y!r} не передбачена', []
            kinds.add(ALLOWED[(x, y)])
    if kinds & {'кома', 'мінус'}:
        m = PURE_NUMBER.fullmatch(raw_cell.strip())     # сира комірка: `0.5` у зворотних лапках — код, не число
        if not m: return 'кома чи мінус у комірці, що не є числом', []
        if m.group(1) in sections: return f'номер підрозділу {m.group(1)} надруковано як дріб', []
    if 'апостроф' in kinds:
        for span in protected(raw_cell):
            if span not in doc: return f'апостроф змінено у вставці коду чи адресі «{span}»', []
    return None, sorted(kinds)


def source_tables(lines, code):
    """Таблиці джерела по порядку: списки номерів рядків, без рядка-роздільника."""
    tables, cur = [], []
    for n, s in enumerate(lines + [''], 1):
        if s.startswith('|') and n not in code:
            cur.append(n)
        elif cur:
            tables.append([k for k in cur
                           if not all(re.fullmatch(r':?-+:?', c) for c in ref_cells(lines[k - 1]))])
            cur = []
    return tables


def body_elements(path):
    """Тіло документа по порядку: [(«P»|«T», текст, [SHA-256 зображень], [ширина зображень, мм],
    орієнтація сторінки)]. Зображення читаються з пакета за r:embed — байти, які побачить читач;
    орієнтація — з розділу, до якого належить елемент (sectPr у абзаці закриває свій розділ)."""
    import hashlib
    doc = Document(path)
    A = '{http://schemas.openxmlformats.org/drawingml/2006/main}blip'
    R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'
    WP = '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}'
    W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

    def orient(sp):
        pg = sp.find(W + 'pgSz')
        if pg is None: return 'portrait'
        if pg.get(W + 'orient') == 'landscape': return 'landscape'
        return 'landscape' if int(pg.get(W + 'w', 0)) > int(pg.get(W + 'h', 0)) else 'portrait'

    body = doc.element.body
    kids = [el for el in body.iterchildren() if el.tag in (W + 'p', W + 'tbl')]
    cur, orients = orient(body.find(W + 'sectPr')), []
    for el in reversed(kids):
        sp = el.find(f'{W}pPr/{W}sectPr')
        if sp is not None: cur = orient(sp)
        orients.append(cur)
    orients.reverse()
    out = []
    for el, o in zip(kids, orients):
        shas, widths = [], []
        for dr in el.iter(WP + 'inline', WP + 'anchor'):
            b, e = dr.find('.//' + A), dr.find(WP + 'extent')
            if b is None: continue
            shas.append(hashlib.sha256(doc.part.related_parts[b.get(R)].blob).hexdigest())
            widths.append(int(e.get('cx')) / 36000 if e is not None else 0.0)
        text = ws(''.join(t.text or '' for t in el.iter(W + 't')))
        out.append(('P' if el.tag == W + 'p' else 'T', text, shas, widths, o))
    return out


def check_figures(fig_dir, docx, lines):
    """Перевірка 7: рисунки на місцях і саме ті. Повертає список порушень."""
    import hashlib
    man = json.load(open(os.path.join(fig_dir, 'manifest.json'), encoding='utf-8'))
    els, bad = body_elements(docx), []
    file_sha = lambda f: hashlib.sha256(open(os.path.join(fig_dir, f), 'rb').read()).hexdigest()
    for num, fig in man.items():
        want = [file_sha(s['file']) for s in fig['sheets']]
        caps = [i for i, (k, t, *_) in enumerate(els)
                if k == 'P' and re.match(rf'^Рис\. {re.escape(num)} \S', t)]
        if len(caps) != 1:
            bad.append(f'рисунок {num}: підписів «Рис. {num} …» у .docx {len(caps)}, має бути 1'); continue
        i = caps[0]
        prev = els[i - 1] if i else ('', '', [], [], '')
        first = want if fig['kind'] == 'images' else want[:1]
        if prev[0] != ('T' if fig['kind'] == 'images' else 'P') or prev[2] != first:
            bad.append(f'рисунок {num}: перед підписом не той рендер (очікувано {len(first)} зображ., '
                       f'знайдено {len(prev[2])}, збіг SHA-256: {prev[2] == first})')
        if fig['kind'] == 'images': continue
        placed = [prev]
        for k, sha in enumerate(want[1:], 2):
            img, lab = (els[i + 2 * k - 3: i + 2 * k - 1] + [None, None])[:2]
            if not img or img[2] != [sha]:
                bad.append(f'рисунок {num}, аркуш {k}: після підпису не той рендер або його немає')
            if not lab or lab[1] != f'Рис. {num}, аркуш {k}':
                bad.append(f'рисунок {num}, аркуш {k}: немає підпису «Рис. {num}, аркуш {k}»')
            placed.append(img)
        # як рисунок стоїть насправді: орієнтація сторінки й кегль найдрібнішого тексту на папері —
        # з ширини зображення в .docx, а не з того, що збирався зробити build.py
        for k, (sheet, el) in enumerate(zip(fig['sheets'], placed), 1):
            if not el or not el[3]: continue
            if el[4] != fig['orientation']:
                bad.append(f'рисунок {num}, аркуш {k}: сторінка {el[4]}, а має бути {fig["orientation"]}')
            pt = sheet['min_font_px'] * el[3][0] / sheet['css_px'][0] / 0.3528
            if pt < MIN_PT - 0.05:
                bad.append(f'рисунок {num}, аркуш {k}: найдрібніший текст на папері {pt:.1f} пт < {MIN_PT}')
    total = sum(len(e[2]) for e in els); expected = sum(len(f['sheets']) for f in man.values())
    if total != expected:
        bad.append(f'зображень у .docx {total}, відрендерено {expected}')
    return bad, man


def source_code_blocks(lines):
    """Блоки коду джерела по порядку: огороджені ``` і діапазони коду з реєстру поза
    огорожами. Порожні рядки всередині огорожі — частина коду, крайні відкинуто; поза
    огорожею (діапазон із реєстру) порожній рядок — межа абзаців markdown, не код."""
    blocks, cur, inside, prev = [], None, False, None
    in_range = lambda n: any(a <= n <= b for a, b in build.CODE_RANGES)
    for n, s in enumerate(lines, 1):
        if inside:
            if s.rstrip().endswith('```'):
                head = s.rstrip()[:-3].rstrip()
                if head.strip(): cur.append(head)
                blocks.append(cur); inside = False
            else:
                cur.append(s)
        elif s.lstrip().startswith('```'):
            rest = s.lstrip()[3:]
            head = rest[len(re.match(r'[A-Za-z]*', rest).group(0)):]
            if s.count('```') >= 2: blocks.append([head.rstrip()[:-3].strip()])
            else: cur, inside = ([head.strip()] if head.strip() else []), True
        elif in_range(n):                    # рядок коду з реєстру поза огорожею
            if prev != n - 1: blocks.append([])
            if s.strip(): blocks[-1].append(s)
            prev = n
    out = []
    for b in blocks:
        while b and not b[0].strip(): b = b[1:]
        while b and not b[-1].strip(): b = b[:-1]
        if b: out.append([ln.rstrip() for ln in b])
    return out


def docx_code_blocks(path):
    """Групи сусідніх абзаців коду в тілі документа: абзац коду — той, де всі runs
    набрано Courier New (порожній рядок коду — один порожній run Courier New)."""
    blocks, cur = [], []
    for p in Document(path).paragraphs:
        if p.runs and all(r.font.name == 'Courier New' for r in p.runs):
            cur.append(p.text.rstrip())
        elif cur:
            blocks.append(cur); cur = []
    if cur: blocks.append(cur)
    return blocks


# ================================================================ перевірка 9: оформлення
WNS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
MNS = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
ITEM = r'(?:\d+|[А-ЯІЇЄҐ])\.\d+'
HEAD_GAP_MM = (6.0, 9.0)                 # методичка, с. 17: 2–3 інтервали між заголовком і текстом


def _body(doc):
    """Тіло документа по порядку: абзаци й таблиці як об'єкти python-docx."""
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    out = []
    for el in doc.element.body.iterchildren():
        if el.tag == WNS + 'p': out.append(Paragraph(el, doc))
        elif el.tag == WNS + 'tbl': out.append(Table(el, doc))
    return out


def _is_p(x): return x.__class__.__name__ == 'Paragraph'
def _align(p): return p.alignment if p.alignment is not None else p.style.paragraph_format.alignment
def _mm(length): return None if length is None else length / 36000
def _has_field(p, name): return any((t.text or '').strip().startswith(name) for t in p._p.iter(WNS + 'instrText'))
def _spaced(p): return any(r._r.find(f'{WNS}rPr/{WNS}spacing') is not None for r in p.runs if r.text.strip())


def check_layout(doc, lines):
    """Повертає список порушень оформлення. `doc` — відкритий Document (самоперевірка передає
    зіпсований у пам'яті), `lines` — рядки full.md: із них береться лише те, скільки чого має бути."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH as AL, WD_LINE_SPACING as LS
    bad, els = [], _body(doc)
    text = lambda x: ws(x.text) if _is_p(x) else ''
    near = lambda a, b: a is not None and abs(a - b) < 0.15

    # аркуш і поля: книжкові розділи 30/10/20/20, альбомний (рисунок) — те саме, повернуте
    for k, sec in enumerate(doc.sections, 1):
        w, h = _mm(sec.page_width), _mm(sec.page_height)
        m = [_mm(x) for x in (sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin)]
        want = ([30, 10, 20, 20], (210, 297)) if h > w else ([20, 20, 30, 10], (297, 210))
        if not (near(w, want[1][0]) and near(h, want[1][1])):
            bad.append(f'розділ документа {k}: аркуш {w:.0f}×{h:.0f} мм, а не А4')
        if not all(near(a, b) for a, b in zip(m, want[0])):
            bad.append(f'розділ документа {k}: поля {[round(x) for x in m]} мм, мають бути {want[0]}')
    n = doc.styles['Normal']; f = n.paragraph_format
    if n.font.name != 'Times New Roman' or n.font.size is None or n.font.size.pt != 14:
        bad.append(f'основний текст: {n.font.name} {n.font.size.pt if n.font.size else "?"} пт, а не Times New Roman 14')
    if f.line_spacing_rule != LS.ONE_POINT_FIVE: bad.append('основний текст: інтервал не півтора')
    if not near(_mm(f.first_line_indent), 12.5): bad.append('основний текст: абзацний відступ не 12,5 мм')
    for name in ('Heading 1', 'Heading 2'):
        gap = _mm(doc.styles[name].paragraph_format.space_after)
        if gap is None or not (HEAD_GAP_MM[0] <= gap <= HEAD_GAP_MM[1]):
            bad.append(f'{name}: до тексту {gap} мм, має бути {HEAD_GAP_MM[0]:.0f}–{HEAD_GAP_MM[1]:.0f} мм')

    # заголовки розділів: «РОЗДІЛ N» без крапки, назва з нового рядка, по центру
    chapters = [x for x in els if _is_p(x) and x.style.name == 'Heading 1' and re.match(r'РОЗДІЛ \d', x.text)]
    want = sum(1 for s in lines if re.match(r'^# РОЗДІЛ \d', s))
    if len(chapters) != want: bad.append(f'заголовків розділів {len(chapters)}, у джерелі {want}')
    for x in chapters:
        if not re.match(r'^РОЗДІЛ \d+\n\S', x.text):
            bad.append(f'«{ws(x.text)[:30]}»: назва розділу не з нового рядка або після номера крапка')
        if _align(x) != AL.CENTER: bad.append(f'«{ws(x.text)[:30]}»: заголовок розділу не по центру')

    # заголовки пунктів: у розрядці, з крапкою, у підбір до тексту, із записом для змісту
    src = [s for s in lines if s.strip()]
    for i, s in enumerate(src):
        m = re.match(r'^### (.+)$', s)
        if not m: continue
        head = m.group(1).rstrip('.') + '.'
        hits = [x for x in els if _is_p(x) and x.text.replace('\u00a0', ' ').startswith(head)]   # типографіка: NBSP перед тире
        if len(hits) != 1:
            bad.append(f'пункт «{head[:40]}»: абзаців із таким початком {len(hits)}'); continue
        x, nxt = hits[0], (src[i + 1] if i + 1 < len(src) else '')
        if not _spaced(x): bad.append(f'пункт «{head[:40]}»: заголовок не в розрядці')
        if not _has_field(x, 'TC'): bad.append(f'пункт «{head[:40]}»: немає запису для змісту')
        if not near(_mm(x.paragraph_format.first_line_indent or f.first_line_indent), 12.5):
            bad.append(f'пункт «{head[:40]}»: не з абзацного відступу')
        prose = not re.match(r'^(#|\||```|Таблиця |Рис\. |Лістинг |\[МІСЦЕ|\d+\.\s|[–—-]\s)', nxt) \
            and not build.RE_FORMULA.match(nxt)
        if prose and len(x.text) <= len(head) + 1:
            bad.append(f'пункт «{head[:40]}»: текст не в підбір до заголовка')

    # таблиці: «Таблиця N» праворуч, назва під ним по центру, далі сама таблиця
    caps = [i for i, x in enumerate(els) if _is_p(x) and re.fullmatch(rf'Таблиця {ITEM}', text(x))]
    want = sum(1 for s in lines if re.fullmatch(rf'Таблиця {ITEM}', s))
    if len(caps) != want: bad.append(f'підписів «Таблиця N» {len(caps)}, у джерелі {want}')
    for i in caps:
        num, title, tbl = els[i], els[i + 1], els[i + 2]
        if _align(num) != AL.RIGHT: bad.append(f'«{text(num)}»: номер таблиці не праворуч')
        if not _is_p(title) or _align(title) != AL.CENTER or not text(title):
            bad.append(f'«{text(num)}»: назва таблиці не по центру під номером')
        if _is_p(tbl): bad.append(f'«{text(num)}»: після назви немає таблиці')

    # продовження таблиці: напис праворуч, з нової сторінки, далі частина з тією самою головкою
    grid_row0 = lambda t: [ws(c.text) for c in t.rows[0].cells]
    last_cap, first_head = None, {}
    for i, x in enumerate(els):
        if _is_p(x):
            m = re.fullmatch(rf'Таблиця ({ITEM})', text(x))
            if m: last_cap = m.group(1)
            m = CONTINUED.fullmatch(text(x))
            if not m: continue
            num, nxt = m.group(1), (els[i + 1] if i + 1 < len(els) else None)
            if num != last_cap: bad.append(f'«{text(x)}» стоїть після таблиці {last_cap}')
            if _align(x) != AL.RIGHT: bad.append(f'«{text(x)}»: напис не праворуч')
            if not x.paragraph_format.page_break_before: bad.append(f'«{text(x)}»: частина не з нової сторінки')
            if nxt is None or _is_p(nxt): bad.append(f'«{text(x)}»: після напису немає таблиці')
            elif grid_row0(nxt) != first_head.get(num): bad.append(f'«{text(x)}»: головка частини не та, що в таблиці')
        elif x.style is not None and x.style.name == 'Table Grid' and last_cap not in first_head:
            first_head[last_cap] = grid_row0(x)
        if not _is_p(x) and x.style is not None and x.style.name == 'Table Grid':
            if any(r._tr.find(f'{WNS}trPr/{WNS}cantSplit') is None for r in x.rows):
                bad.append(f'таблиця {last_cap}: рядок може розірватися між сторінками'); 

    # рисунки: «Рис. N Назва» під рисунком по центру, без крапки в кінці
    figs = [x for x in els if _is_p(x) and re.match(rf'Рис\. {ITEM} \S', text(x))]
    want = sum(1 for s in lines if re.match(rf'^Рис\. {ITEM} \S', s))
    if len(figs) != want: bad.append(f'підписів «Рис. N» {len(figs)}, у джерелі {want}')
    for x in figs:
        if _align(x) != AL.CENTER: bad.append(f'«{text(x)[:30]}»: підпис рисунка не по центру')
        if text(x).endswith('.'): bad.append(f'«{text(x)[:30]}»: крапка в кінці підпису рисунка')
    if any(_is_p(x) and re.match(r'Рисунок \S+ [–—-]', text(x)) for x in els):
        bad.append('лишився підпис у формі «Рисунок N – Назва»')

    # формули: вільний рядок над і під кожною (між двома поспіль — один)
    is_formula = lambda x: not _is_p(x) and x._tbl.find('.//' + MNS + 'oMath') is not None
    blank = lambda x: _is_p(x) and not x.text.strip() and not x._p.findall('.//' + WNS + 'drawing')
    formulas = [i for i, x in enumerate(els) if is_formula(x)]
    if len(formulas) != len({m.group(2) for s in lines for m in [build.RE_FORMULA.match(s)] if m and m.group(2) in FORMULA_TEX}):
        bad.append(f'формул у документі {len(formulas)}, у джерелі інша кількість')
    for i in formulas:
        num = re.search(r'\((\d+\.\d+)\)', ' '.join(c.text for r in els[i].rows for c in r.cells))
        name = f'формула ({num.group(1)})' if num else 'формула'
        if not blank(els[i - 1]): bad.append(f'{name}: над нею немає вільного рядка')
        if i + 1 >= len(els) or not blank(els[i + 1]): bad.append(f'{name}: під нею немає вільного рядка')

    # додатки: «Додаток А» праворуч, назва під ним по центру, кожен (крім першого) з нової сторінки
    apps = [i for i, x in enumerate(els) if _is_p(x) and re.fullmatch(r'Додаток [А-ЯІЇЄҐ]', text(x))]
    want = sum(1 for s in lines if re.match(r'^## Додаток [А-ЯІЇЄҐ] ', s))
    if len(apps) != want: bad.append(f'заголовків «Додаток X» {len(apps)}, у джерелі {want}')
    for k, i in enumerate(apps):
        x, title = els[i], els[i + 1]
        if _align(x) != AL.RIGHT: bad.append(f'«{text(x)}»: не праворуч')
        if not _is_p(title) or _align(title) != AL.CENTER: bad.append(f'«{text(x)}»: назва додатка не по центру')
        if not _has_field(x, 'TC'): bad.append(f'«{text(x)}»: немає запису для змісту')
        if k and not x.paragraph_format.page_break_before: bad.append(f'«{text(x)}»: не з нової сторінки')

    # перелік умовних позначень: дві колонки, рядків стільки, скільки в джерелі
    heads = [i for i, x in enumerate(els) if _is_p(x) and x.style.name == 'Heading 1' and text(x).startswith('ПЕРЕЛІК УМОВНИХ')]
    if heads:
        a = next(i for i, s in enumerate(lines) if s.startswith('# ПЕРЕЛІК УМОВНИХ'))
        b = next(i for i in range(a + 1, len(lines)) if lines[i].startswith('# '))
        want = sum(1 for s in lines[a + 1:b] if build.LEGEND_ITEM.match(s))
        end = next(i for i in range(heads[0] + 1, len(els)) if _is_p(els[i]) and els[i].style.name == 'Heading 1')
        tabs = [x for x in els[heads[0]:end] if not _is_p(x)]
        rows = sum(len(t.rows) for t in tabs)
        if not tabs or any(len(t.columns) != 2 for t in tabs) or rows != want:
            bad.append(f'перелік умовних позначень: рядків у двох колонках {rows}, у джерелі {want}')

    # номер сторінки: поле PAGE у верхньому колонтитулі праворуч
    # номер сторінки: картка, титульний лист і завдання — без номера; зміст — сторінка 3,
    # номер — поле PAGE у верхньому колонтитулі праворуч (методичка, с. 10 і 18)
    if len(doc.sections) < 2:
        bad.append('немає окремого розділу документа для ненумерованих перших сторінок'); return bad
    if any(_has_field(p, 'PAGE') for p in doc.sections[0].header.paragraphs):
        bad.append('титульний лист і завдання мають номер сторінки')
    hp = doc.sections[1].header.paragraphs[0]
    if not _has_field(hp, 'PAGE') or _align(hp) != AL.RIGHT:
        bad.append('номер сторінки: немає поля PAGE вгорі праворуч')
    start = doc.sections[1]._sectPr.find(WNS + 'pgNumType')
    if start is None or start.get(WNS + 'start') != '3':
        bad.append('зміст не є сторінкою 3: нумерований розділ не починається з трьох')
    first_numbered = next((ws(x.text) for x in _section_start(doc, 1) if _is_p(x) and x.text.strip()), '')
    if first_numbered != 'ЗМІСТ':
        bad.append(f'перша нумерована сторінка починається з «{first_numbered[:30]}», а не зі змісту')
    bad += check_typography(doc)
    return bad


# Типографіка (власник, 2026-10-26) — перевіряється в готовому .docx власними правилами, не кодом
# build.py: у тексті поза кодом і адресами не лишилося звичайного пробілу перед « –», між числом і
# одиницею чи всередині «в. п.», звичайного дефіса в слові через дефіс і символу U+2011 (у Times New
# Roman немає його гліфа — дефіс має бути елементом <w:noBreakHyphen/>). І навпаки: нерозривні
# пробіли й дефіси в документі є — перевірка, що може сказати лише «нічого не знайшла», нічого не доводить.
TY_ADDR = re.compile(r'\S*(?:https?://|www\.|@|\.md\b|\.py\b|\.json\b|\.sql\b|doi\.org)\S*')
TY_BAD = (
    ('звичайний пробіл перед тире', re.compile(r'\S –(\s|$)')),
    ('звичайний пробіл між числом і одиницею', re.compile(
        r'\d (%|в\.\s?п\.|с\.|мс|хв|год|ГБ|МБ|мм|пт|px|тис\.|млн|дол\.|євро)(?!\w)|\d с(?=[\s,;:)]|$)')),
    ('звичайний пробіл у «в. п.»', re.compile(r'\bв\. п\.')),
    ('звичайний дефіс у слові через дефіс', re.compile(r'[^\W_]-[^\W_]')),
    ('символ U+2011 замість елемента Word', re.compile('‑')),
)


def check_typography(doc):
    """Текст кожного абзацу збирається з його прогонів (код — пробілом, нерозривний дефіс — знаком
    U+2027), і правила застосовуються до абзацу: дефіс, що стоїть окремим вузлом між двома словами,
    теж помітний."""
    bad, nbsp, nbh = [], 0, 0
    for p in doc.element.body.iter(WNS + 'p'):
        parts = []
        for r in p.iter(WNS + 'r'):
            rf = r.find(f'{WNS}rPr/{WNS}rFonts')
            if rf is not None and 'Courier' in (rf.get(WNS + 'ascii') or ''):
                parts.append(' '); continue
            for el in r:
                if el.tag == WNS + 't': parts.append(el.text or '')
                elif el.tag == WNS + 'noBreakHyphen': parts.append('\u2027'); nbh += 1
        txt = TY_ADDR.sub(' ', ''.join(parts))
        nbsp += txt.count('\u00a0')
        for what, rx in TY_BAD:
            m = rx.search(txt)
            if m: bad.append(f'типографіка: {what} — «{txt[max(0, m.start() - 25):m.end() + 15]}»')
    if not nbsp: bad.append('типографіка: у документі немає жодного нерозривного пробілу')
    # інлайнова математика: індекси — справжні (vertAlign), і сирого «σ_shape», «A_m(x)» у тексті
    # поза кодом немає: окрема літера, підкреслення, індекс. Ідентифікатори на кшталт calendar_events
    # (слово перед підкресленням) правилом не зачіпаються.
    subs = sum(1 for v in doc.element.body.iter(WNS + 'vertAlign') if v.get(WNS + 'val') == 'subscript')
    if not subs: bad.append('математика: у документі немає жодного справжнього нижнього індексу')
    for p in doc.element.body.iter(WNS + 'p'):
        txt = ''.join(t.text or '' for r in p.iter(WNS + 'r')
                      if not ((rf := r.find(f'{WNS}rPr/{WNS}rFonts')) is not None and 'Courier' in (rf.get(WNS + 'ascii') or ''))
                      for t in r.findall(WNS + 't'))
        m = re.search(r'(?<![\w])[^\W\d_][\u0300-\u036f]*_[\w{(]|\$[^$]+\$', TY_ADDR.sub(' ', txt))
        if m: bad.append(f'математика: сира розмітка в тексті — «{txt[max(0, m.start() - 20):m.end() + 15]}»')
    if not nbh: bad.append('типографіка: у документі немає жодного нерозривного дефіса')
    return bad


def _section_start(doc, k):
    """Елементи тіла від початку k-го розділу документа (розділ закриває sectPr в останньому
    його абзаці)."""
    els, seen = [], 0
    for x in _body(doc):
        if seen >= k: els.append(x)
        if _is_p(x) and x._p.find(f'{WNS}pPr/{WNS}sectPr') is not None: seen += 1
    return els


def layout_mutations():
    """Самоперевірка перевірки 9: (що зіпсовано, функція-псувальник, фрагмент очікуваного повідомлення)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
    from docx.shared import Mm, Pt

    def first(doc, pred): return next(x for x in _body(doc) if pred(x))
    def cap(doc): return first(doc, lambda x: _is_p(x) and re.fullmatch(rf'Таблиця {ITEM}', ws(x.text)))
    def drop(el): el.getparent().remove(el)

    def no_break(doc):
        h = first(doc, lambda x: _is_p(x) and x.style.name == 'Heading 1' and x.text.startswith('РОЗДІЛ'))
        for br in list(h._p.iter(WNS + 'br')): br.getparent().replace(br, parse_xml_t(' '))
    def unspace(doc):
        h = first(doc, lambda x: _is_p(x) and _spaced(x) and _has_field(x, 'TC'))   # пункт, а не «ЗАВДАННЯ» форми
        for sp in list(h._p.iter(WNS + 'spacing')):
            if sp.getparent().tag == WNS + 'rPr': drop(sp)
    def no_blank_above(doc):
        els = _body(doc); i = next(i for i, x in enumerate(els) if not _is_p(x) and x._tbl.find('.//' + MNS + 'oMath') is not None)
        drop(els[i - 1]._p)
    def appendix_left(doc):
        first(doc, lambda x: _is_p(x) and re.fullmatch(r'Додаток [А-ЯІЇЄҐ]', ws(x.text))).alignment = AL.LEFT
    def legend_row(doc):
        els = _body(doc); i = next(i for i, x in enumerate(els) if _is_p(x) and x.text.startswith('ПЕРЕЛІК УМОВНИХ'))
        t = next(x for x in els[i:] if not _is_p(x)); drop(t.rows[0]._tr)
    def plain_space(doc):
        t = next(t for t in doc.element.body.iter(WNS + 't') if '\u00a0–' in (t.text or ''))
        t.text = t.text.replace('\u00a0–', ' –', 1)
    def plain_hyphen(doc):
        h = next(doc.element.body.iter(WNS + 'noBreakHyphen'))
        h.getparent().replace(h, parse_xml_t('-'))
    def raw_math(doc):
        v = next(v for v in doc.element.body.iter(WNS + 'vertAlign') if v.get(WNS + 'val') == 'subscript')
        r = v.getparent().getparent(); t = r.find(WNS + 't'); t.text = '_' + t.text
        v.getparent().remove(v)
    def spoil_head(doc):
        els = _body(doc); i = next(i for i, x in enumerate(els) if _is_p(x) and CONTINUED.fullmatch(ws(x.text)))
        els[i + 1].rows[0].cells[0].paragraphs[0].add_run('?')
    return [
        ('поле 25 мм замість 30', lambda d: setattr(d.sections[0], 'left_margin', Mm(25)), 'поля'),
        ('кегль 12 замість 14', lambda d: setattr(d.styles['Normal'].font, 'size', Pt(12)), 'Times New Roman 14'),
        ('6 пт між заголовком підрозділу і текстом', lambda d: setattr(d.styles['Heading 2'].paragraph_format, 'space_after', Pt(6)), 'Heading 2'),
        ('назва розділу в одному рядку з «РОЗДІЛ N»', no_break, 'не з нового рядка'),
        ('заголовок пункту без розрядки', unspace, 'не в розрядці'),
        ('номер таблиці ліворуч', lambda d: setattr(cap(d), 'alignment', AL.LEFT), 'не праворуч'),
        ('формула без вільного рядка над нею', no_blank_above, 'над нею немає вільного рядка'),
        ('«Додаток А» ліворуч', appendix_left, 'не праворуч'),
        ('перелік без одного рядка', legend_row, 'перелік умовних позначень'),
        ('«Продовження таблиці» ліворуч',
         lambda d: setattr(first(d, lambda x: _is_p(x) and CONTINUED.fullmatch(ws(x.text))), 'alignment', AL.LEFT),
         'напис не праворуч'),
        ('частина таблиці з чужою головкою', spoil_head, 'головка частини не та'),
        ('звичайний пробіл перед тире', plain_space, 'пробіл перед тире'),
        ('звичайний дефіс замість нерозривного', plain_hyphen, 'дефіс у слові'),
        ('індекс набрано сирим «_»', raw_math, 'сира розмітка'),
        ('зміст не на третій сторінці',
         lambda d: d.sections[1]._sectPr.find(WNS + 'pgNumType').set(WNS + 'start', '2'), 'не є сторінкою 3'),
    ]


def parse_xml_t(txt):
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    return parse_xml('<w:t %s xml:space="preserve">%s</w:t>' % (nsdecls('w'), txt))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--self-test', action='store_true',
                    help='перевірка 9 на навмисно зіпсованому документі: кожне правило має спрацювати')
    ap.add_argument('--src', required=True); ap.add_argument('--docx', required=True)
    ap.add_argument('--brief'); ap.add_argument('--registry', default=os.path.join(HERE, 'registry.json'))
    ap.add_argument('--figures', help='тека рендерів із manifest.json: перевірка 7')
    ap.add_argument('--private', help='titulka.json: перевірка 8 (приватні імена лише в .docx)')
    ap.add_argument('--referat', help='реферат окремим .docx: його рядки шукаються в ньому, а не в роботі')
    ap.add_argument('--pages', type=int, help='обсяг, з яким складено документ (build.py --pages)')
    ap.add_argument('--body-pages', type=int, help='обсяг основного тексту, з яким складено документ (build.py --body-pages)')
    a = ap.parse_args()

    if a.self_test:
        lines = open(a.src, encoding='utf-8').read().split('\n')
        clean = check_layout(Document(a.docx), lines)
        failed = len(clean)
        for t in clean: print(f'   документ до псування вже порушує: {t}')
        for what, spoil, needle in layout_mutations():
            doc = Document(a.docx); spoil(doc)
            got = check_layout(doc, lines)
            ok = any(needle in g for g in got)
            print(f'   {"ok  " if ok else "НЕ СПРАЦЮВАЛА"} {what}' + ('' if ok else f' — {got[:2]}'))
            failed += not ok
        print(f'самоперевірка оформлення: {len(layout_mutations())} псувань, не помічено {failed}')
        return 1 if failed else 0

    build.load(a.src, a.brief)
    private = dict(build.load_private(a.private, a.pages, a.body_pages))   # дані, як реєстр; перевіряє їх перевірка 8
    if a.registry: build.bind_registry(a.registry)
    LINES = build.LINES
    body, nmath = docx_text(a.docx)
    # реферат не вшивається в роботу (методичка, с. 11): його рядки мають бути в окремому файлі
    # і НЕ мають бути в роботі
    _, ua, en, _toc = build._front_lines()
    referat_lines = set(range(ua[0], _toc))
    referat_body = docx_text(a.referat)[0] if a.referat else None
    code, sections = code_lines(LINES), heading_numbers(LINES)
    filled = set(json.load(open(os.path.join(a.figures, 'manifest.json'), encoding='utf-8'))) if a.figures else set()
    failed = False

    # 1, 3
    missing, formulas, lost_words, apostrophes = [], [], [], 0
    for n in range(1, len(LINES) + 1):
        s = raw(n)
        if not s.strip(): continue
        ph = build.RE_FIG_PH.match(s)
        if ph and ph.group(1) in filled:                         # замінено рисунком — перевірка 7
            if strip_md(s) in body: missing.append((n, 'ПЛЕЙСХОЛДЕР ЛИШИВСЯ поруч із рисунком: ' + strip_md(s)[:60]))
            continue
        fm = build.RE_FORMULA.match(s)
        if fm and fm.group(2) in FORMULA_TEX:
            num = fm.group(2); formulas.append(num)
            for word in re.findall(r'[А-ЯІЇЄҐа-яіїєґ]{3,}', fm.group(1)):
                if word not in FORMULA_TEX[num]: lost_words.append((n, num, word))
            # розділовий знак у кінці формули (кома між формулами поспіль — методичка, с. 20)
            # набрано двічі: у рядку full.md і в LaTeX; вони мають збігатися
            end_src = fm.group(1).rstrip()[-1:]
            for tex in [FORMULA_TEX[num]] + ([FORMULA_SPLIT[num][-1]] if num in FORMULA_SPLIT else []):
                end_tex = tex.rstrip()[-1:]
                if (end_src in ',;.') != (end_tex in ',;.') or (end_src in ',;.' and end_src != end_tex):
                    lost_words.append((n, num, f'кінцевий знак: у full.md «{end_src}», у formulas.py «{end_tex}»'))
            continue
        if s.startswith('|') and n not in code: continue            # таблиці — перевірка 2
        if n in code:
            t = ws(re.sub(r'^\s*```[a-z]*\s*|\s*```\s*$', '', s))
            if len(t) >= 3 and t not in body: missing.append((n, t[:90]))
            continue
        t = strip_md(s)
        if n in referat_lines:
            if referat_body is None: continue
            if len(t) >= 40 and (t in body or strip_md(curl(s)) in body):
                missing.append((n, 'РЯДОК РЕФЕРАТУ ВШИТО В РОБОТУ: ' + t[:60]))
            if len(t) < 3 or t in referat_body: continue
            if strip_md(curl(s)) in referat_body: apostrophes += 1
            else: missing.append((n, 'немає в рефераті: ' + t[:70]))
            continue
        if len(t) < 3 or t in body: continue
        if strip_md(curl(s)) in body: apostrophes += 1
        else: missing.append((n, t[:90]))

    # 2
    src_tabs, doc_tabs = source_tables(LINES, code), docx_tables(a.docx)
    cells_bad, normalized = [], []
    if len(src_tabs) != len(doc_tabs):
        cells_bad.append((0, f'таблиць у джерелі {len(src_tabs)}, у .docx {len(doc_tabs)}'))
    for rows, drows in zip(src_tabs, doc_tabs):
        if len(rows) != len(drows):
            cells_bad.append((rows[0], f'рядків {len(rows)} ↔ {len(drows)}')); continue
        for n, dcells in zip(rows, drows):
            scells = ref_cells(raw(n))
            if len(scells) != len(dcells):
                cells_bad.append((n, f'комірок {len(scells)} ↔ {len(dcells)}')); continue
            for j, (sc, dc) in enumerate(zip(scells, dcells), 1):
                why, kinds = judge(sc, strip_md(sc), dc, sections)
                if why: cells_bad.append((n, f'комірка {j}: {why}: «{strip_md(sc)[:40]}» → «{dc[:40]}»'))
                elif kinds: normalized.append((n, j, strip_md(sc), dc, kinds))

    extra = sum(len(FORMULA_SPLIT[k]) - 1 for k in FORMULA_SPLIT if k in formulas)
    ok_f = nmath == len(formulas) + extra
    print(f'рядків джерела перевірено: {len(LINES)}')
    print(f'рядків тексту, не знайдених у .docx: {len(missing)}'
          + (f' (рядків, де апостроф став ’: {apostrophes})' if apostrophes else ''))
    for n, t in missing[:20]: print(f'   ряд {n}: {t}')
    print(f'таблиць: джерело {len(src_tabs)}, .docx {len(doc_tabs)}; комірок, що не збіглися: {len(cells_bad)}; '
          f'нормалізованих комірок: {len(normalized)}')
    for n, t in cells_bad[:20]: print(f'   ряд {n}: {t}')
    for n, j, x, y, k in normalized[:20]: print(f'   (нормалізовано) ряд {n}, комірка {j}: «{x}» → «{y}» ({", ".join(k)})')
    print(f'формул у джерелі: {len(formulas)}; обʼєктів OMML: {nmath}'
          + (f' (з них {extra} — другі рядки розбитих)' if extra else '') + ('  ✔' if ok_f else '  ✘'))
    for n, num, w in lost_words: print(f'   РОЗБІЖНІСТЬ У ФОРМУЛІ ({num}) між full.md і formulas.py, ряд {n}: «{w}»')
    failed |= bool(missing or cells_bad or lost_words or not ok_f)

    # 4
    if a.brief:
        src_txt = fold(ws(strip_md('\n'.join(LINES))))
        kv = build.KEEP_VERBATIM
        absent_src = [k for k in kv if fold(strip_md(k)) not in src_txt]
        printed = fold(body + ' ' + (referat_body or ''))     # робота і реферат: рядок має бути в одному з них
        not_in_docx = [k for k in kv if k not in absent_src and fold(strip_md(k)) not in printed]
        exempt = [k for k in not_in_docx if k in DOCX_EXEMPT]
        bad = [k for k in not_in_docx if k not in DOCX_EXEMPT]
        print(f'keep-verbatim: {len(kv)} рядків; у джерелі немає {len(absent_src)}; '
              f'у .docx втрачено {len(bad)}; виняток «структура рівняння» {len(exempt)}')
        for k in absent_src: print(f'   НЕМАЄ В ДЖЕРЕЛІ: {k}')
        for k in bad: print(f'   ВТРАЧЕНО В .docx: {k}')
        for k in exempt: print(f'   (виняток) {k}')
        failed |= bool(absent_src or bad)

    # 5
    ins_bad = [t for lst in build.INSERT_AFTER.values() for t in lst
               if strip_md(t) not in body and strip_md(t).replace("'", '’') not in body]
    print(f'вставок із реєстру застосовано: {sum(len(v) for v in build.INSERT_AFTER.values())}; '
          f'не знайдено в .docx: {len(ins_bad)}')
    failed |= bool(ins_bad)

    # 8
    if a.private:
        public = open(a.src, encoding='utf-8').read()           # сирий full.md, як лежить у репозиторії
        raw_values = [str(v).strip() for k, v in json.load(open(a.private, encoding='utf-8')).items()
                      if k in NAME_FIELDS and str(v or '').strip()]   # кафедра й спеціальність — не таємниця
        norm = lambda t: re.sub(r'\s+', ' ', t).casefold()
        priv_bad = [f'ПРИВАТНЕ ЗНАЧЕННЯ У full.md: «{v}»' for v in raw_values if norm(v) in norm(public)]
        # прізвище — перше слово з великої літери, довше за ініціал, — цілим словом без огляду на
        # регістр: «Прізвище І. І.» чи саме прізвище в тексті — теж витік. Ім'я окремо не
        # шукається: імена бувають звичайними словами («Надія», «Віра»), а без прізвища не
        # вказують на людину.
        names = [str(v) for k, v in json.load(open(a.private, encoding='utf-8')).items()
                 if k in NAME_FIELDS and str(v or '').strip()]
        surnames = [next((t for t in re.findall(r"[^\W\d_][\w'’-]{2,}", v) if t[0].isupper()), None) for v in names]
        for tok in sorted({t for t in surnames if t}):
            if re.search(rf'(?<![\w-]){re.escape(tok)}(?![\w-])', public, re.IGNORECASE):
                priv_bad.append(f'ПРІЗВИЩЕ У full.md: «{tok}»')
        both = body + ' ' + (referat_body or '')            # титулка — у роботі, реферат — окремим файлом
        for _, ph, _ in build.PRIVATE_FIELDS:
            if ph not in public:
                priv_bad.append(f'плейсхолдера «{ph}» у full.md немає')
                continue
            if ph in private:
                if ph in both: priv_bad.append(f'плейсхолдер «{ph}» лишився в .docx, хоча значення задано')
                if private[ph] not in both: priv_bad.append(f'значення для «{ph}» немає в .docx')
            elif ph not in both:
                priv_bad.append(f'незаповнений плейсхолдер «{ph}» зник із .docx')
        for ph in build.PAGE_PLACEHOLDERS:
            if a.pages and (ph in both or str(a.pages) not in both):
                priv_bad.append(f'обсяг {a.pages} с. не підставлено замість «{ph}»')
        # обсяг основного тексту — власне правило: число стоїть саме в реченні ВСТУПУ про основний текст
        if a.body_pages and (build.BODY_PLACEHOLDER in body
                             or not re.search(rf'основний текст\s–\s{a.body_pages}\sс\.', body)):
            priv_bad.append(f'обсяг основного тексту {a.body_pages} с. не підставлено замість «{build.BODY_PLACEHOLDER}»')
        filled = sum(1 for _, ph, _ in build.PRIVATE_FIELDS if ph in private)
        print(f'дані титулки: заповнено {filled} з {len(build.PRIVATE_FIELDS)}; порушень: {len(priv_bad)}')
        for t in priv_bad: print(f'   {t}')
        failed |= bool(priv_bad)

    # 7
    if a.figures:
        fig_bad, man = check_figures(a.figures, a.docx, LINES)
        print(f'рисунків відрендерено: {len(man)} (аркушів і знімків {sum(len(f["sheets"]) for f in man.values())}); '
              f'порушень: {len(fig_bad)}')
        for t in fig_bad: print(f'   {t}')
        failed |= bool(fig_bad)

    # 6
    sb, db = source_code_blocks(LINES), docx_code_blocks(a.docx)
    code_bad = [] if len(sb) == len(db) else [f'блоків коду в джерелі {len(sb)}, у .docx {len(db)}']
    for k, (x, y) in enumerate(zip(sb, db), 1):
        if x != y:
            j = next((i for i, (u, v) in enumerate(zip(x, y)) if u != v), min(len(x), len(y)))
            code_bad.append(f'блок {k}, рядок {j + 1}: «{(x[j] if j < len(x) else "⟨кінець⟩")[:50]}» ↔ '
                            f'«{(y[j] if j < len(y) else "⟨кінець⟩")[:50]}»')
    print(f'блоків коду: джерело {len(sb)}, .docx {len(db)} '
          f'(рядків {sum(map(len, sb))}, з них порожніх {sum(ln == "" for b in sb for ln in b)}); '
          f'розбіжностей: {len(code_bad)}')
    for t in code_bad[:20]: print(f'   {t}')
    failed |= bool(code_bad)

    # 9
    layout_bad = check_layout(Document(a.docx), open(a.src, encoding='utf-8').read().split('\n'))
    print(f'оформлення за методичкою: порушень {len(layout_bad)}')
    for t in layout_bad[:30]: print(f'   {t}')
    failed |= bool(layout_bad)

    print('РЕЗУЛЬТАТ:', 'ПОМИЛКИ' if failed else 'OK')
    return 1 if failed else 0


if __name__ == '__main__':
    sys.exit(main())
