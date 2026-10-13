#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Складальник магістерської роботи: full.md → .docx за методичними рекомендаціями кафедри (2026).

    python3 build.py --src ../text/full.md --brief ../formatter-brief.md \\
                     --registry registry.json --out $DOCS/build/hourwell.docx --report $DOCS/build/report.json \\
                     --figures $DOCS/build/figures --private $DOCS/private/titulka.json

Головне правило: жоден рядок прози не набирається руками — текст читається з full.md.
Руками зроблено лише те, що не має текстового вигляду: формули (formulas.py) і реєстр
погоджених правок (registry.json). Кожне рішення, закладене в код, описане в README.md.

Залежності: python-docx, lxml, pandoc (для перетворення LaTeX формул в OMML).
"""
import argparse, json, os, re, subprocess, sys, tempfile

from docx import Document
from docx.shared import Pt, Mm, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from formulas import FORMULA_TEX, FORMULA_SPLIT, FORMULA_PT   # noqa: E402

C, R, J, LEFT = (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT,
                 WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT)
TW = 170                                  # ширина набору, мм: 210 − 30 − 10
SYMBOL_FONT = 'Segoe UI Symbol'           # для знаків, яких немає в Times New Roman
SYMBOLS = '◐◑◒◓●○◆◇★☆✓✗▲▼'

# ---------------------------------------------------------------- стан
LINES = []                 # рядки джерела (1-based через L())
KEEP_VERBATIM = []         # рядки брифа, які мають пережити складання
OVERRIDE = {}              # номер рядка → {find, replace, why}
INSERT_AFTER = {}          # номер рядка → [текст абзацу]
CODE_RANGES = []           # (перший, останній) — блоки коду поза огорожами
REPORT = {'placeholders': [], 'notes': [], 'figures': [], 'tables': [],
          'listings': [], 'formulas': []}
NORMALIZED = {'кома': [], 'мінус': [], 'апостроф': []}
APOSTROPHE_KV = set()
_FENCE_LINES = None
_OMML = {}


# ================================================================ дані титулки
# Титулка — форма кафедри (методичка, Додаток В). Її дані живуть поза репозиторієм, у
# MagisterDocs/private/titulka.json (власник, 2026-09-29; кафедра й спеціальність — туди ж,
# 2026-10-13): у full.md (публічний репозиторій) лишаються плейсхолдери, а значення підставляються
# лише під час складання .docx. Порожнє значення — плейсхолдер лишається. Третя колонка каже, чи це
# ім'я: імена не мають з'явитися в full.md за жодних умов (verify.py, перевірка 8); кафедра й
# спеціальність — не таємниця, вони просто зберігаються в одному місці з рештою даних форми.
PRIVATE_FIELDS = (                        # (ключ у titulka.json, плейсхолдер у full.md, ім'я?)
    ('здобувач', '[ПІБ здобувача]', True),
    ('здобувач_англійською', '[Student name]', True),
    ('здобувач_на_титулці', '[ПІБ здобувача для титулки]', True),
    ('керівник', '[ПІБ керівника]', True),
    ('керівник_ступінь_і_звання', '[науковий ступінь і вчене звання керівника]', False),
    ('гарант', '[ПІБ гаранта]', True),
    ('гарант_ступінь_і_звання', '[науковий ступінь і вчене звання гаранта]', False),
    ('декан', '[ПІБ декана]', True),
    ('завідувач_кафедри', '[ПІБ завідувача кафедри]', True),
    ('кафедра', '[назва кафедри]', False),
    ('спеціальність', '[код і найменування спеціальності]', False),
    ('спеціальність_англійською', '[speciality code and name]', False),
)
PAGE_PLACEHOLDERS = ('[кількість сторінок]', '[number of pages]')   # обсяг роботи: з готового PDF
PRIVATE = {}                              # плейсхолдер → значення (лише заповнені)
DATA = {}                                 # ключ titulka.json → значення: для приватних частин (картка, завдання)


def load_private(path, pages=None):
    """Читає titulka.json і підставляє заповнені значення в LINES. Повертає словник
    плейсхолдер → значення. Невідомий ключ у файлі валить складання (друкарська помилка в
    ключі інакше мовчки лишила б плейсхолдер). `pages` — кількість сторінок роботи, виміряна
    на готовому PDF попереднього проходу (run_all.sh): число в тексті не набирається руками."""
    global LINES
    PRIVATE.clear(); DATA.clear()
    if pages:
        for ph in PAGE_PLACEHOLDERS: PRIVATE[ph] = str(pages)
    if path:
        data = json.load(open(path, encoding='utf-8'))
        known = {k for k, _, _ in PRIVATE_FIELDS}
        unknown = [k for k in data if not k.startswith('_') and k not in known]
        if unknown:
            raise SystemExit(f'titulka.json: невідомі ключі {unknown}; відомі — {sorted(known)}')
        for key, ph, _ in PRIVATE_FIELDS:
            val = str(data.get(key) or '').strip()
            DATA[key] = val
            if val: PRIVATE[ph] = val
    lost = [ph for ph in PRIVATE if not any(ph in ln for ln in LINES)]
    if lost:                              # інакше заповнене значення мовчки не потрапило б нікуди
        raise SystemExit(f'titulka.json: значення задано, але плейсхолдера {lost} у full.md немає')
    LINES = [_sub_private(ln) for ln in LINES]
    return PRIVATE


def _sub_private(line):
    # довші плейсхолдери першими: «[ПІБ здобувача]» є початком «[ПІБ здобувача для титулки]» лише
    # без дужки, але порядок не має залежати від такого збігу
    for ph in sorted(PRIVATE, key=len, reverse=True):
        line = line.replace(ph, PRIVATE[ph])
    return line


def load(src, brief=None):
    """Завантажує джерело й бриф. Скидає весь стан — модуль можна перевикористати."""
    global LINES, KEEP_VERBATIM, _FENCE_LINES
    LINES = open(src, encoding='utf-8').read().split('\n')
    KEEP_VERBATIM = []
    if brief:
        m = re.search(r'```keep-verbatim\n(.*?)```', open(brief, encoding='utf-8').read(), re.S)
        if m: KEEP_VERBATIM = m.group(1).strip().split('\n')
    OVERRIDE.clear(); INSERT_AFTER.clear(); CODE_RANGES.clear(); APOSTROPHE_KV.clear()
    for v in REPORT.values(): v.clear()
    for v in NORMALIZED.values(): v.clear()
    _FENCE_LINES = None


# ================================================================ нормалізація
# Три масові заміни, погоджені з автором. Кожна — правилом, із запобіжниками,
# які валять складання, якщо заміна зачепила те, що чіпати не можна.

def _is_code_line(n):
    """Рядок усередині огородженого блоку ``` (разом із самими огорожами)
    або в одному з CODE_RANGES."""
    global _FENCE_LINES
    if _FENCE_LINES is None:
        _FENCE_LINES, inside = set(), False
        for i, ln in enumerate(LINES, 1):
            if ln.lstrip().startswith('```'):
                _FENCE_LINES.add(i)
                inside = not ln.rstrip().endswith('```') or ln.count('```') == 1 and not inside
                continue
            if inside: _FENCE_LINES.add(i)
            if inside and ln.rstrip().endswith('```'): inside = False
    return n in _FENCE_LINES or any(a <= n <= b for a, b in CODE_RANGES)


NUMERIC_CELL = re.compile(r'^[+\-−]?\d+(\.\d+)?(\s*±\s*\d+(\.\d+)?)?$')
SECTION_COLUMN = re.compile(r'(?<!\w)(?:[Пп]ід)?[Рр]озділ')   # «Розділ реалізації», «Підрозділ»
CODE_SPAN = re.compile(r'`[^`]*`')
URLISH = re.compile(r'\S*(?:https?://|www\.|@|\.md|\.py|\.json|\.sql)\S*')


def _cell_spans(s):
    """Межі комірок рядка таблиці за GFM: зворотна скісна екранує наступний символ,
    тож \\| — риска всередині комірки. Повертає [(a, b)]: s[a:b] — вміст комірки
    без рисок-роздільників."""
    bars, i = [], 0
    while i < len(s):
        if s[i] == '\\': i += 2; continue
        if s[i] == '|': bars.append(i)
        i += 1
    spans = [(a + 1, b) for a, b in zip(bars, bars[1:])]
    if bars and s[bars[-1] + 1:].strip(): spans.append((bars[-1] + 1, len(s)))
    return spans


def _table_header(n):
    """Рядок заголовка таблиці, якій належить рядок n."""
    h = n
    while h > 1 and LINES[h - 2].startswith('|'): h -= 1
    return LINES[h - 1]


def _decimal_comma(n, s):
    """Десяткова крапка → кома і ASCII-дефіс → мінус. Лише в суто числових
    комірках таблиць: у прозі «підрозділ 5.2» під це правило не потрапляє.

    Стовпці, чий заголовок називає розділ чи підрозділ, правило теж оминає: `3.7`
    там — номер підрозділу, а не дріб. Табл. Е.1, «Розділ реалізації» для FR-42:
    конвеєр, відновлений 2026-09-28, друкував його як «3,7».

    ВИНЯТОК, ЯКИЙ НЕ Є ПОМИЛКОЮ: англомовна анотація лишається з десятковою
    крапкою — «+0.4 pp», «1.4–1.9 pp». В англійському тексті крапка є нормою,
    а кома читається як роздільник тисяч, тож «1,4 pp» англомовний рецензент
    прочитає як «одна тисяча чотириста». Правило сюди не дістає, бо анотація —
    проза, а не числова комірка таблиці. Цей самий виняток записаний у брифі (§7);
    перевірка «одна десяткова позначка в документі», якщо колись з'явиться, має
    вносити англомовну анотацію до винятків саме з цієї причини."""
    if not s.startswith('|') or _is_code_line(n): return s
    head = _table_header(n)
    skip = {j for j, (a, b) in enumerate(_cell_spans(head)) if SECTION_COLUMN.search(head[a:b])}
    out = s
    for j, (a, b) in enumerate(_cell_spans(s)):
        cell = s[a:b].strip()
        if j in skip or not (cell and NUMERIC_CELL.match(cell)): continue
        new = cell.replace('.', ',')
        if new.startswith('-'): new = '−' + new[1:]
        if new != cell:
            out = out[:a] + s[a:b].replace(cell, new) + out[b:]     # довжина та сама — межі тримаються
            if ',' in new and '.' in cell: NORMALIZED['кома'].append((n, cell, new))
            if new.startswith('−') and cell.startswith('-'): NORMALIZED['мінус'].append((n, cell, new))
    return out


def _apostrophe(n, s):
    """Прямий апостроф → ’. Код, адреси й вставки в зворотних лапках недоторканні:
    у f'y_{t.id}' апостроф — синтаксис, а не пунктуація."""
    if "'" not in s: return s
    if _is_code_line(n): return s
    holes = []

    def stash(m):
        holes.append(m.group(0)); return '\x00%d\x00' % (len(holes) - 1)
    tmp = CODE_SPAN.sub(stash, s)
    tmp = URLISH.sub(stash, tmp)
    if "'" in tmp:
        NORMALIZED['апостроф'].append((n, tmp.count("'")))
        tmp = tmp.replace("'", '\u2019')
    for i, h in enumerate(holes):
        tmp = tmp.replace('\x00%d\x00' % i, h)
    return tmp


def _normalize(n, s):
    before = s
    s = _decimal_comma(n, s)
    s = _apostrophe(n, s)
    if s != before:
        fold = lambda t: t.replace("'", '\u2019')          # бриф §7: чекери згортають обидва апострофи
        for k in KEEP_VERBATIM:
            if fold(k) in fold(before) and fold(k) not in fold(s):
                raise SystemExit(f'НОРМАЛІЗАЦІЯ ЗАЧЕПИЛА ЗАХИЩЕНИЙ РЯДОК, ряд {n}: «{k}»')
            if k in before and k not in s and fold(k) in fold(s):
                APOSTROPHE_KV.add(k)
        if len(before) != len(s):
            raise SystemExit(f'НОРМАЛІЗАЦІЯ ЗМІНИЛА ДОВЖИНУ РЯДКА {n} — це не заміна символів')
        for a, b in zip(before, s):
            if a != b and (a, b) not in (('.', ','), ('-', '−'), ("'", '\u2019')):
                raise SystemExit(f'НОРМАЛІЗАЦІЯ ЗМІНИЛА {a!r} → {b!r} у ряд. {n} — не передбачено правилами')
    return s


def L(n):
    """Рядок джерела: погоджені override-и + нормалізація. Єдиний шлях, яким текст
    потрапляє в документ."""
    if n in OVERRIDE:
        o = OVERRIDE[n]; raw = LINES[n-1]
        if o['replace'] in raw:
            note = f'ряд {n}: правку вже внесено в джерело, override не застосовано'
        elif o['find'] in raw:
            raw = raw.replace(o['find'], o['replace'], 1)
            note = f'ряд {n}: ЗАМІНА за погодженням — {o["why"]}'
        else:
            raise SystemExit(f'OVERRIDE РЯД. {n}: не знайдено ані «{o["find"][:40]}», ані «{o["replace"][:40]}»')
        if note not in REPORT['notes']: REPORT['notes'].append(note)
        return _normalize(n, raw)
    return _normalize(n, LINES[n-1])


# ================================================================ якорі й реєстр
# Реєстр правок прив'язаний до тексту, а не до номерів рядків: після
# перегенерації full.md номери зсуваються, а якір лишається. Ненайдений або
# неоднозначний якір ВАЛИТЬ складання — мовчки зібраний документ без правки
# гірший за відсутній.

def anchor(text):
    hits = [i for i, ln in enumerate(LINES, 1) if text in ln]
    if not hits:
        raise SystemExit(f'ЯКІР НЕ ЗНАЙДЕНО: «{text[:70]}»')
    if len(hits) > 1:
        raise SystemExit(f'ЯКІР НЕОДНОЗНАЧНИЙ ({len(hits)} збігів): «{text[:70]}» → рядки {hits}')
    return hits[0]


def bind_registry(path):
    reg = json.load(open(path, encoding='utf-8'))
    src_all = '\n'.join(LINES)
    for o in reg.get('overrides', []):
        OVERRIDE[anchor(o['anchor'])] = o
    for ins in reg.get('inserts', []):
        n = anchor(ins['after_anchor'])
        if ins.get('present_if') and ins['present_if'] in src_all:
            REPORT['notes'].append(f'вставку після ряд. {n} не застосовано: '
                                   f'«{ins["present_if"][:40]}…» уже є в джерелі')
            continue
        INSERT_AFTER.setdefault(n, []).append(ins['text'])
    for cb in reg.get('code_blocks', []):
        h = anchor(cb['heading_anchor'])
        # від першого «{» після заголовка до останнього «}» перед наступним розділом
        z0 = next(i for i in range(h + 1, h + 6) if LINES[i-1].strip() == '{')
        nxt = next(i for i in range(h + 1, len(LINES) + 1) if LINES[i-1].startswith('## '))
        z1 = next(i for i in range(nxt - 1, h, -1) if LINES[i-1].strip() == '}')
        fences = [i for i in range(z0, z1 + 1) if LINES[i-1].lstrip().startswith('```')]
        if fences:
            REPORT['notes'].append(f'ДЕФЕКТ ДЖЕРЕЛА, {cb["heading_anchor"]}: закривальна огорожа стоїть у '
                                   f'ряд. {fences[0]}, а блок триває до ряд. {z1} — складальник узяв увесь блок як код')
        CODE_RANGES.append((z0, z1))
    n_anchors = len(reg.get('overrides', [])) + len(reg.get('inserts', [])) + len(reg.get('code_blocks', []))
    REPORT['notes'].append(f'реєстр привʼязано за {n_anchors} якорями, усі знайдено')


# ================================================================ стилі
def new_doc():
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin, s.right_margin = Mm(30), Mm(10)
    s.top_margin, s.bottom_margin = Mm(20), Mm(20)

    def force(style, name='Times New Roman', size=14, bold=None):
        style.font.name = name; style.font.size = Pt(size)
        if bold is not None: style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None: rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
        for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), name)

    n = doc.styles['Normal']; force(n, bold=False)
    p = n.paragraph_format
    p.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.first_line_indent = Cm(1.25); p.alignment = J
    p.space_before = Pt(0); p.space_after = Pt(0)

    # Методичка, с. 17: між заголовком і текстом — 2–3 інтервали (6–9 мм); 18 пт = 6,35 мм.
    # Заголовок пункту під це правило не підпадає: його друкують у розрядці в підбір до тексту
    # (runin_heading) — це початок абзацу, а не стиль заголовка.
    for name, before, after, align, ind, brk, bold in (
            ('Heading 1', 0, 18, C, 0, True, True),
            ('Heading 2', 18, 18, LEFT, 1.25, False, True)):
        st = doc.styles[name]; force(st, size=14, bold=bold)
        f = st.paragraph_format
        f.space_before, f.space_after = Pt(before), Pt(after)
        f.alignment, f.first_line_indent = align, Cm(ind)
        f.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        f.keep_with_next = True; f.page_break_before = brk
    # Рядки змісту: без власних стилів Word і LibreOffice беруть кожен свої відступи рівнів.
    # Підрозділи, пункти й додатки — з одним відступом: запис додатка LibreOffice ставить на
    # рівень нижче, ніж Word, і різний відступ рівнів 2 і 3 розвів би їх у PDF.
    from docx.enum.style import WD_STYLE_TYPE
    for lvl, left in ((1, 0), (2, 0.5), (3, 0.5)):
        try: st = doc.styles['toc %d' % lvl]
        except KeyError: st = doc.styles.add_style('toc %d' % lvl, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles['Normal']; force(st, bold=False)
        f = st.paragraph_format
        f.first_line_indent = Cm(0); f.left_indent = Cm(left); f.alignment = LEFT
        f.space_before = f.space_after = Pt(0)
        f.tab_stops.add_tab_stop(Mm(TW), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    return doc


def numbered_section(doc, start):
    """Розділ документа з номером сторінки вгорі праворуч, нумерація з `start`. Усе перед ним —
    картка, титульний лист і завдання — номерів не має (методичка, с. 10)."""
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    brk = doc.element.body.findall(qn('w:p'))[-1]              # порожній абзац із розривом — у попередній
    prev = brk.getprevious()
    if prev is not None and prev.tag == qn('w:p') and prev.find(qn('w:pPr') + '/' + qn('w:sectPr')) is None:
        ppr = prev.find(qn('w:pPr'))
        if ppr is None: ppr = OxmlElement('w:pPr'); prev.insert(0, ppr)
        ppr.append(brk.find(qn('w:pPr') + '/' + qn('w:sectPr'))); brk.getparent().remove(brk)
    s.header.is_linked_to_previous = False; s.footer.is_linked_to_previous = False
    page_number(s.header.paragraphs[0])
    for pg in s._sectPr.findall(qn('w:pgNumType')):            # успадковане від перших сторінок («з нуля»)
        s._sectPr.remove(pg)
    s._sectPr.append(parse_xml('<w:pgNumType %s w:start="%d"/>' % (nsdecls('w'), start)))
    return s


def page_number(p):
    """Поле PAGE у колонтитулі, праворуч."""
    p.alignment = R
    p.paragraph_format.first_line_indent = Cm(0)
    for kind, val in (('begin', None), ('instr', ' PAGE '), ('end', None)):
        r = p.add_run()
        if kind == 'instr':
            e = OxmlElement('w:instrText'); e.text = val
        else:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind)
        r._r.append(e)


# ================================================================ inline-розмітка
TOKEN = re.compile(r'(\*\*.+?\*\*|`[^`]+`|(?<![\w\\])_[^_\n]+_(?![\w]))', re.S)


def _is_prose_italic(inner):
    """Курсивом стає лише проза. Ідентифікатор, шлях чи адреса — ніколи,
    навіть якщо підкреслення випадково склалися в пару."""
    core = inner.rstrip(':;,.!? ')
    if ' ' in core: return True                        # проза — курсив доречний
    if '/' in core or '\\' in core: return False       # шлях
    if re.search(r'\w_\w', core): return False         # ідентифікатор
    return core.isalpha() and core != ''


def add_literal(p, text, size=None, mono=None):
    """Текст без жодної інтерпретації розмітки."""
    r = p.add_run(text)
    if size: r.font.size = Pt(size)
    if mono:
        r.font.name = mono
        r._element.rPr.rFonts.set(qn('w:ascii'), mono)
        r._element.rPr.rFonts.set(qn('w:hAnsi'), mono)
    return p


def add_runs(p, text, size=None, bold_all=False, italic_all=False):
    """Розбір вкладеної розмітки. Вставки в зворотних лапках усередині курсиву
    чи жирного обробляються теж — інакше лапки лишаються в тексті."""
    for part in TOKEN.split(text):
        if not part: continue
        bold, italic, mono, txt = bold_all, italic_all, False, part
        if part.startswith('**') and part.endswith('**'):
            add_runs(p, part[2:-2], size=size, bold_all=True, italic_all=italic_all); continue
        elif part.startswith('`') and part.endswith('`'):
            mono, txt = True, part[1:-1]
        elif part.startswith('_') and part.endswith('_'):
            if _is_prose_italic(part[1:-1]):
                add_runs(p, part[1:-1], size=size, bold_all=bold_all, italic_all=True); continue
            txt = part
        r = p.add_run(txt)
        if bold: r.bold = True
        if italic: r.italic = True
        if size: r.font.size = Pt(size)
        if mono:
            r.font.name = 'Courier New'
            r._element.rPr.rFonts.set(qn('w:ascii'), 'Courier New')
            r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Courier New')
    # знаки, яких немає в Times New Roman, переносимо в символьний шрифт
    for r in list(p.runs):
        if any(ch in SYMBOLS for ch in r.text) and r.font.name != 'Courier New':
            pieces = re.split('([' + SYMBOLS + '])', r.text)
            if len(pieces) > 1:
                r.text = pieces[0]
                anchor_run = r
                for piece in pieces[1:]:
                    if not piece: continue
                    nr = p.add_run(piece)
                    nr.bold, nr.italic, nr.font.size = r.bold, r.italic, r.font.size
                    if piece in SYMBOLS:
                        nr.font.name = SYMBOL_FONT
                        nr._element.rPr.rFonts.set(qn('w:ascii'), SYMBOL_FONT)
                        nr._element.rPr.rFonts.set(qn('w:hAnsi'), SYMBOL_FONT)
                    anchor_run._element.addnext(nr._element); anchor_run = nr
    return p


def para(doc, text='', *, align=None, indent=None, spacing=None, size=None, bold=False,
         mono=None, before=None, after=None, style=None, keep=False, literal=False):
    p = doc.add_paragraph(style=style)
    if literal and not mono:
        add_literal(p, text, size=size)
    elif mono:
        r = p.add_run(text); r.font.name = mono
        r._element.rPr.rFonts.set(qn('w:ascii'), mono); r._element.rPr.rFonts.set(qn('w:hAnsi'), mono)
        if size: r.font.size = Pt(size)
    elif text:
        add_runs(p, text, size=size, bold_all=bold)
    f = p.paragraph_format
    if align is not None: f.alignment = align
    if indent is not None: f.first_line_indent = Cm(indent)
    if spacing == 1: f.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if before is not None: f.space_before = Pt(before)
    if after is not None: f.space_after = Pt(after)
    if keep: f.keep_with_next = True
    return p


def toc_field(doc):
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0)
    for kind, val in (('begin', None), ('instr', r' TOC \o "1-3" \h \z \u \f '), ('separate', None)):
        r = p.add_run()
        if kind == 'instr':
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = val
        else:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind)
        r._r.append(e)
    p.add_run('Зміст збирається полем Word: виділити все (Ctrl+A) і натиснути F9.')
    r = p.add_run(); e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r._r.append(e)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


RE_CHAPTER = re.compile(r'^(РОЗДІЛ \d+) (.+)$')
RE_APPENDIX = re.compile(r'^(Додаток [А-ЯІЇЄҐ]) (.+)$')


def struct_head(doc, text, in_toc=True, brk=True):
    """Структурний заголовок. in_toc=False — не потрапляє в автозміст (сам ЗМІСТ).
    Методичка, с. 18: «РОЗДІЛ 1» — без крапки, назва розділу — з нового рядка. Це один абзац із
    розривом рядка: у змісті він стає одним записом «РОЗДІЛ 1 НАЗВА»."""
    if in_toc:
        p = doc.add_paragraph(style='Heading 1')
        m = RE_CHAPTER.match(text)
        if m:
            add_runs(p, m.group(1)); p.runs[-1].add_break(WD_BREAK.LINE); add_runs(p, m.group(2))
        else:
            add_runs(p, text)
        p.paragraph_format.page_break_before = brk
    else:
        if brk: page_break(doc)
        p = para(doc, text, align=C, indent=0, bold=True, after=18, keep=True)
    return p


def field(p, instr):
    """Поле Word без видимого результату (TC — запис для змісту)."""
    for kind in ('begin', 'instr', 'end'):
        r = p.add_run()
        if kind == 'instr':
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = instr
        else:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind)
        r._r.append(e)


def appendix_head(doc, label, title, brk):
    """Методичка, с. 22: «Додаток А» — праворуч, під ним назва симетрично до тексту, кожен
    додаток — з нової сторінки. Два абзаци з різним вирівнюванням не можуть бути одним
    заголовком, тож у зміст додаток іде записом TC (поле Word), одним рядком «Додаток А Назва»."""
    p = para(doc, label, align=R, indent=0, keep=True, literal=True)
    p.paragraph_format.page_break_before = brk
    field(p, ' TC "%s %s" \\l 2 ' % (label, title.replace('"', '\u201d')))
    t = para(doc, '', align=C, indent=0, after=18, keep=True)
    add_runs(t, title, bold_all=True)
    return p


SPACED = 30                               # розрядка заголовка пункту, двадцяті частки пункту: 1,5 пт
TC_LEVEL = 2                              # рівень записів TC у змісті — див. toc_entry


def toc_entry(p, text):
    """Запис для змісту (поле TC) на початку абзацу `p`. Рівень 2 і для додатків, і для пунктів:
    Word ставить запис на вказаний рівень, а LibreOffice — на один нижче (перевірено на 26.8.0:
    «\\l 3» у зміст PDF не потрапляє зовсім), тож рівень 2 дає запис в обох, а однаковий відступ
    стилів «toc 2» і «toc 3» робить різницю невидимою."""
    field(p, ' TC "%s" \\l %d ' % (text.replace('"', '\u201d'), TC_LEVEL))


def appendix_head(doc, label, title, brk):
    """Методичка, с. 22: «Додаток А» — праворуч, під ним назва симетрично до тексту, кожен
    додаток — з нової сторінки. Два абзаци з різним вирівнюванням не можуть бути одним
    заголовком, тож у зміст додаток іде записом TC, одним рядком «Додаток А Назва»."""
    p = para(doc, '', align=R, indent=0, keep=True)
    p.paragraph_format.page_break_before = brk
    toc_entry(p, label + ' ' + title)
    add_literal(p, label)
    t = para(doc, '', align=C, indent=0, after=18, keep=True)
    add_runs(t, title, bold_all=True)
    return p


def runin_heading(doc, text):
    """Заголовок пункту (методичка, с. 17): з абзацного відступу, у розрядці, у підбір до тексту,
    з крапкою в кінці. Це початок звичайного абзацу, а не окремий абзац-заголовок: текст пункту
    дописується в той самий абзац (build_body), і так рядок виходить однаковим у Word і в PDF.
    Прихований знак абзацу (роздільник стилів Word) LibreOffice верстає рамкою, яка на заголовку,
    довшому за рядок, розриває текст. У зміст пункт іде записом TC."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.keep_with_next = True
    head = text if text.rstrip().endswith('.') else text.rstrip() + '.'
    toc_entry(p, head)
    first = len(p.runs)
    add_runs(p, head)
    for r in p.runs[first:]:
        sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), str(SPACED)); r._r.get_or_add_rPr().append(sp)
    return p


def boxed(doc, text):
    """Плейсхолдер у пунктирній рамці: видно на сторінці, що тут бракує вставки."""
    p = para(doc, text, indent=0, align=LEFT, size=12, spacing=1, before=6, after=6)
    p._p.get_or_add_pPr().append(parse_xml(
        '<w:pBdr %s>' % nsdecls('w') +
        ''.join('<w:%s w:val="dashed" w:sz="6" w:space="6" w:color="808080"/>' % s
                for s in ('top', 'left', 'bottom', 'right')) + '</w:pBdr>'))
    return p


# ================================================================ класифікатор
RE_H = re.compile(r'^(#{1,3})\s+(.*)$')
RE_FORMULA = re.compile(r'^(.*?)\s{2,}\((\d+\.\d+)\)\s*$')
RE_FIG_PH = re.compile(r'^(?:\*\*)?\[МІСЦЕ ДЛЯ РИСУНКА ([^\]]+)\](?:\*\*)?')
RE_FILL_PH = re.compile(r'\[МІСЦЕ ДЛЯ ЗАПОВНЕННЯ:')
RE_FIG_CAP = re.compile(r'^Рис\. ((?:\d+|[А-ЯІЇЄҐ])\.\d+) \S')        # «Рис. 1.2 Назва» (методичка, с. 19)
RE_TAB_CAP = re.compile(r'^Таблиця ((?:\d+|[А-ЯІЇЄҐ])\.\d+)$')        # номер; назва — наступним абзацом
RE_LST_CAP = re.compile(r'^Лістинг ([\w.]+)\s+[—–-]')
RE_NUM_ITEM = re.compile(r'^\d+\.\s')
RE_BULLET = re.compile(r'^[–—-]\s')
RE_CONSTRAINT = re.compile(r'^\((C\d)\)\s+')
CYR = re.compile(r'^[А-ЯІЇЄҐа-яіїєґ«»(]')


# ================================================================ таблиці
SPLITS = {}                               # номер таблиці → [рядків тіла в кожній частині] (layout.py)
NEW_PAGE = set()                          # таблиці, що починаються з нової сторінки (під головкою лишався один рядок)
CUR_TABLE = [None]                        # номер таблиці з останнього підпису «Таблиця N»
CONTINUED = 'Продовження таблиці'


def load_splits(path):
    """План поділу таблиць із layout.py (README, «Продовження таблиці»). Немає файла — таблиці цілі."""
    SPLITS.clear(); NEW_PAGE.clear(); CUR_TABLE[0] = None
    if path and os.path.exists(path):
        for num, v in json.load(open(path, encoding='utf-8')).items():
            if v.get('parts'): SPLITS[num] = v['parts']
            if v.get('before'): NEW_PAGE.add(num)


def split_cells(line):
    """Комірки рядка таблиці за правилами GFM: зворотна скісна екранує наступний
    символ, тож \\| лишає риску всередині комірки (|зміщення|/SE, |A_m(x)|, |приріст|),
    а риска без скісної перед нею — роздільник. Після розбору \\| стає «|».

    Раніше тут стояло «зшивання» комірок, що закінчуються скісною. Воно давало той самий
    результат для правильного екранування, але мовчки лагодило й зламане (скісна, пробіли,
    риска) — тобто ховало дефект джерела. Тепер зламане екранування валить складання."""
    inner = line.strip()
    if inner.startswith('|'): inner = inner[1:]
    if inner.endswith('|') and not inner.endswith('\\|'): inner = inner[:-1]
    cells, cur, i = [], '', 0
    while i < len(inner):
        ch = inner[i]
        if ch == '\\' and i + 1 < len(inner):
            cur += inner[i:i+2]; i += 2; continue
        if ch == '|':
            cells.append(cur); cur = ''
        else:
            cur += ch
        i += 1
    cells.append(cur)
    out = []
    for c in cells:
        c = c.strip()
        if c.endswith('\\') and not c.endswith('\\\\'):
            raise SystemExit(f'ЗЛАМАНЕ ЕКРАНУВАННЯ РИСКИ в рядку таблиці «{line[:60]}…»: '
                             f'комірка «{c[-30:]}» закінчується скісною, відокремленою від риски')
        out.append(c.replace('\\|', '|'))
    return out


def md_table_rows(start):
    rows, i = [], start
    while i <= len(LINES) and L(i).startswith('|'):
        rows.append(split_cells(L(i))); i += 1
    if len(rows) > 1 and set(rows[1][0]) <= set('-: '): del rows[1]
    ncol = len(rows[0])
    bad = [k for k, r in enumerate(rows[1:], 1) if len(r) != ncol]
    if bad:
        raise SystemExit(f'ТАБЛИЦЯ В РЯД. {start}: після зшивання кількість комірок не збігається '
                         f'(заголовок {ncol}, рядки {bad})')
    return rows, i


def emit_table(doc, start):
    rows, nxt = md_table_rows(start)
    header, body = rows[0], rows[1:]
    empty_hdr = [j for j, h in enumerate(rows[0]) if not h.strip()]
    if empty_hdr:
        REPORT['notes'].append(f'УВАГА, таблиця в ряд. {start}: порожні заголовки стовпців '
                               f'{empty_hdr} — ознака розщеплення рисками, перевірити вручну')
    ncol = len(header)
    size = 12 if ncol <= 7 else (11 if ncol <= 9 else 10)
    clean = lambda s: re.sub(r'\*\*|`|_', '', s or '')
    col = lambda j: [clean(header[j])] + [clean(r[j] if j < len(r) else '') for r in body]

    # ширина стовпця: спершу вміщаємо найдовше слово (щоб слова не рвалися),
    # решту простору ділимо пропорційно до обсягу тексту; нижче 10 пт не опускаємось
    def layout(size):
        cw = {12: 2.3, 11: 2.1, 10: 1.9, 9: 1.7}[size]
        minw, mass = [], []
        for j in range(ncol):
            cells = col(j)
            longest = max((len(w) for c in cells for w in c.split()), default=3)
            minw.append(min(60, longest * cw + 4))
            mass.append(sum(len(c) for c in cells) + 1)
        extra = TW - sum(minw)
        if extra < 0: return None
        tot = sum(mass)
        return [round(m + extra * w / tot) for m, w in zip(minw, mass)]

    widths = layout(size)
    while widths is None and size > 10:
        size -= 1; widths = layout(size)
    if widths is None:                                  # не вміщається — стискаємо пропорційно
        base = [len(max(col(j), key=len)) or 3 for j in range(ncol)]
        widths = [max(12, round(TW * b / sum(base))) for b in base]
    over = sum(widths) - TW
    while over > 0:
        k = widths.index(max(widths)); widths[k] -= 1; over -= 1
    while over < 0:
        k = widths.index(min(widths)); widths[k] += 1; over += 1

    def cell(c, txt, bold=False, center=False):
        c.text = ''
        p = c.paragraphs[0]
        f = p.paragraph_format
        f.first_line_indent = Cm(0); f.line_spacing_rule = WD_LINE_SPACING.SINGLE
        f.space_before = Pt(0); f.space_after = Pt(0)
        p.alignment = C if center else LEFT
        add_runs(p, txt.replace('\\|', '|'), size=size, bold_all=bold)

    numeric = [all(re.match(r'^[+−\-–—±◐~]?[\d,.\s%×()]*$', (r[j] if j < len(r) else '') or '')
                   for r in body) for j in range(ncol)]

    def part(rows_):
        """Одна частина таблиці: головка й рядки. Рядок не рветься між сторінками (cantSplit) —
        інакше його половини на двох сторінках не можна ані порахувати, ані підписати."""
        t = doc.add_table(rows=1 + len(rows_), cols=ncol)
        t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
        t._tbl.tblPr.append(parse_xml('<w:tblLayout %s w:type="fixed"/>' % nsdecls('w')))
        t._tbl.tblPr.append(parse_xml('<w:tblW %s w:w="%d" w:type="dxa"/>' % (nsdecls('w'), int(sum(widths) * 56.7))))
        for j, w in enumerate(widths):
            t.columns[j].width = Mm(w)                    # tblGrid — саме його читає Word
            for row in t.rows: row.cells[j].width = Mm(w)
        for j, txt in enumerate(header): cell(t.rows[0].cells[j], txt, bold=True, center=True)
        for i, row in enumerate(rows_, 1):
            for j in range(ncol):
                cell(t.rows[i].cells[j], row[j] if j < len(row) else '', center=numeric[j])
        t.rows[0]._tr.get_or_add_trPr().append(parse_xml('<w:tblHeader %s/>' % nsdecls('w')))
        for row in t.rows:
            row._tr.get_or_add_trPr().append(parse_xml('<w:cantSplit %s/>' % nsdecls('w')))
        for c in t.rows[0].cells:                         # головка не лишається внизу сторінки сама
            for p in c.paragraphs: p.paragraph_format.keep_with_next = True

    # Методичка, с. 20: над кожною наступною частиною таблиці, перенесеної на іншу сторінку, —
    # «Продовження таблиці N». Де таблиця переходить на нову сторінку, знає лише верстка: поділ
    # приходить ззовні (SPLITS — з layout.py, що читає готовий PDF попереднього проходу), і кожна
    # наступна частина починається з нової сторінки цим написом і повтореною головкою.
    sizes = SPLITS.get(CUR_TABLE[0]) or [len(body)]
    if sum(sizes) != len(body) or min(sizes) < 1:
        raise SystemExit(f'ТАБЛИЦЯ {CUR_TABLE[0]}: поділ {sizes} не відповідає {len(body)} рядкам — '
                         f'застарілий файл поділу; почати з проходу без нього')
    at = 0
    for k, size_k in enumerate(sizes):
        if k:
            p = para(doc, f'{CONTINUED} {CUR_TABLE[0]}', indent=0, align=R, spacing=1, after=6, keep=True, literal=True)
            p.paragraph_format.page_break_before = True
        part(body[at:at + size_k]); at += size_k
    para(doc, '', after=6, spacing=1)
    REPORT['tables'].append(f'ряд {start}: {ncol}×{len(body)}, кегль {size}, ширини {widths}'
                            + (f', частин {len(sizes)}: {sizes}' if len(sizes) > 1 else ''))
    return nxt


# ================================================================ код
def mono_size(lines):
    """Найбільший кегль, за якого найдовший рядок ще вміщається в 170 мм.
    Courier New моноширинний і для кирилиці. Округлення вниз — щоб гарантовано."""
    longest = max((len(ln) for ln in lines), default=1)
    pt = 481.9 / (0.6 * max(longest, 1))
    return max(8.0, min(9.5, int(pt * 2) / 2))


def emit_fenced(doc, start):
    """Блок в огорожі ```. Вміст іде дослівно, моноширинним, без розбору розмітки.
    Закривальна огорожа може стояти окремим рядком або в кінці змістового.
    Порожні рядки всередині блоку — теж код: вони відділяють у SQL Додатка Г і в
    лістингу 4.1 одну групу інструкцій від іншої. До 2026-09-28 вони випадали."""
    first = L(start).lstrip()
    rest = first[3:]
    lang = rest.split()[0] if rest.split() and rest.split()[0].isalpha() else ''
    body = rest[len(lang):]
    lines, i = [], start + 1
    if body.rstrip().endswith('```'):                    # відкрито й закрито в тому ж рядку
        lines.append(body.rstrip()[:-3].strip())
    else:
        lines.append(body.strip())
        while i <= len(LINES):
            cur = L(i)
            if cur.rstrip().endswith('```'):
                lines.append(cur.rstrip()[:-3].rstrip()); i += 1; break
            if cur.lstrip().startswith('```'):
                i += 1; break
            lines.append(cur)
            i += 1
    while lines and not lines[0].strip(): lines.pop(0)       # порожнє від рядка-огорожі
    while lines and not lines[-1].strip(): lines.pop()
    sz = mono_size(lines)
    for ln in lines:
        para(doc, ln, indent=0, spacing=1, mono='Courier New', size=sz, align=LEFT)
    para(doc, '', after=6, spacing=1)
    flat = any(len(ln) > 200 for ln in lines)
    REPORT['listings'].append(
        f'ряд {start}: блок ```{lang} — {len(lines)} рядк(ів), кегль {sz}'
        + (', У ДЖЕРЕЛІ ЗЛИТИЙ В ОДИН РЯДОК — переноси втрачено' if flat else ''))
    return i


def emit_listing(doc, start):
    """Лістинг без огорожі (старий формат джерела): до першого рядка, що
    починається кирилицею, або заголовка."""
    i = start
    while i <= len(LINES):
        s = L(i)
        if s == '': i += 1; continue
        if RE_H.match(s) or CYR.match(s): break
        para(doc, s, indent=0, spacing=1, mono='Courier New', size=9.5, align=LEFT)
        i += 1
    para(doc, '', after=6, spacing=1)
    return i


# ================================================================ формули
M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'


def compile_formulas():
    """LaTeX → OMML через pandoc. Одним викликом для всіх формул."""
    if _OMML: return
    tex = dict(FORMULA_TEX)
    for num, parts in FORMULA_SPLIT.items():
        for i, t in enumerate(parts, 1):
            tex['%s#%d' % (num, i)] = t
    keys = list(tex)
    with tempfile.TemporaryDirectory() as d:
        md, dx = os.path.join(d, 'f.md'), os.path.join(d, 'f.docx')
        open(md, 'w', encoding='utf-8').write('\n\n'.join('$%s$' % tex[k] for k in keys))
        subprocess.run(['pandoc', md, '-o', dx], check=True)
        xml = subprocess.run(['unzip', '-p', dx, 'word/document.xml'],
                             capture_output=True, check=True).stdout.decode('utf-8')
    found = re.findall(r'<m:oMath>.*?</m:oMath>', xml, re.S)
    if len(found) != len(keys):
        raise SystemExit(f'PANDOC: скомпільовано {len(found)} формул із {len(keys)}')
    for k, frag in zip(keys, found):
        _OMML[k] = frag.replace('<m:t>-</m:t>', '<m:t>\u2212</m:t>')      # мінус, не дефіс


def omml_el(key, pt):
    compile_formulas()
    frag = _OMML[key].replace('<m:oMath>',
        '<m:oMath xmlns:m="%s" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">' % M_NS, 1)
    el = parse_xml(frag)
    half = str(int(pt * 2))
    for r in el.iter('{%s}r' % M_NS):
        rpr = OxmlElement('w:rPr')
        for tag in ('w:sz', 'w:szCs'):
            e = OxmlElement(tag); e.set(qn('w:val'), half); rpr.append(e)
        mrpr = r.find('{%s}rPr' % M_NS)
        r.insert(list(r).index(mrpr) + 1 if mrpr is not None else 0, rpr)
    return el


def _formula_row(t, r, key, pt, label, number):
    cells = t.rows[r].cells
    ci = 0
    if label is not None:
        c = cells[0]; c.text = ''
        p = c.paragraphs[0]; p.paragraph_format.first_line_indent = Cm(0); p.alignment = LEFT
        if label: p.add_run(label)
        ci = 1
    c = cells[ci]; c.text = ''
    p = c.paragraphs[0]; p.paragraph_format.first_line_indent = Cm(0); p.alignment = C
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
    # oMathPara: Word розкладає вираз як окремий рядок і, якщо він ширший за комірку,
    # переносить його на бінарному операторі, а номер лишається у своїй комірці
    omp = OxmlElement('m:oMathPara')
    ompr = OxmlElement('m:oMathParaPr'); jc = OxmlElement('m:jc'); jc.set(qn('m:val'), 'center')
    ompr.append(jc); omp.append(ompr)
    omp.append(omml_el(key, pt)); p._p.append(omp)
    c = cells[ci + 1]; c.text = ''
    p = c.paragraphs[0]; p.paragraph_format.first_line_indent = Cm(0); p.alignment = R
    if number: p.add_run('(%s)' % number)
    c.vertical_alignment = 1


def emit_formula(doc, n):
    """Формула — невидима таблиця: [мітка обмеження] | формула | номер.
    Номер не може «з'їхати» на наступний рядок чи вліво: він у своїй комірці."""
    s = L(n)
    m = RE_FORMULA.match(s)
    body, num = m.group(1), m.group(2)
    label = None
    cm = RE_CONSTRAINT.match(body)
    if cm: label = '(%s)' % cm.group(1)
    pt = FORMULA_PT.get(num, 12 if len(body) > 74 else 14)
    parts = FORMULA_SPLIT.get(num)
    keys = ['%s#%d' % (num, i) for i in (1, 2)] if parts else [num]
    ncol = 3 if label is not None else 2
    t = doc.add_table(rows=len(keys), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    t._tbl.tblPr.append(parse_xml('<w:tblLayout %s w:type="fixed"/>' % nsdecls('w')))
    widths = [14, TW - 14 - 18, 18] if label is not None else [TW - 18, 18]
    for j, w in enumerate(widths):
        t.columns[j].width = Mm(w)
        for row in t.rows: row.cells[j].width = Mm(w)
    for r, key in enumerate(keys):
        last = (r == len(keys) - 1)
        _formula_row(t, r, key, 14 if parts else pt,
                     (label if r == 0 else '') if label is not None else None,
                     num if last else None)
    REPORT['formulas'].append(f'({num}) ряд {n}, кегль {14 if parts else pt}'
                              + (', розбито на два рядки' if parts else '')
                              + (f', мітка {label}' if label else ''))
    return t


# ================================================================ складання
def _emit_inserts(doc, n):
    for extra in INSERT_AFTER.get(n, []):
        para(doc, _normalize(n, extra))
        REPORT['notes'].append(f'ВСТАВКА після ряд. {n}: {extra[:50]}…')


# ---------------------------------------------------------------- форми: таблиці без рамок
def grid(doc, widths, rows=1):
    """Таблиця без рамок фіксованої ширини — каркас форми (титулка, приватні частини)."""
    t = doc.add_table(rows=rows, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    t._tbl.tblPr.append(parse_xml('<w:tblLayout %s w:type="fixed"/>' % nsdecls('w')))
    for j, w in enumerate(widths):
        t.columns[j].width = Mm(w)
        for row in t.rows: row.cells[j].width = Mm(w)
    return t


def cell_lines(c, lines):
    """Рядки комірки: [(текст, {align, bold, size, underline, rule})]. `rule` — лінія для підпису
    під порожнім абзацом (межа абзацу, а не символи підкреслення)."""
    c.text = ''
    for k, (txt, o) in enumerate(lines):
        p = c.paragraphs[0] if k == 0 else c.add_paragraph()
        f = p.paragraph_format
        f.first_line_indent = Cm(0); f.line_spacing_rule = WD_LINE_SPACING.SINGLE
        f.space_before = Pt(o.get('before', 0)); f.space_after = Pt(o.get('after', 0))
        p.alignment = o.get('align', LEFT)
        if txt:
            r = p.add_run(txt); r.bold = o.get('bold', False); r.underline = o.get('underline', False)
            if o.get('size'): r.font.size = Pt(o['size'])
        if o.get('rule'):
            p._p.get_or_add_pPr().append(parse_xml(
                '<w:pBdr %s><w:bottom w:val="single" w:sz="6" w:space="1" w:color="000000"/></w:pBdr>' % nsdecls('w')))


TITLE_SHAPE = {2: 'ПОГОДЖЕНО', 5: 'ДОПУСКАЄТЬСЯ ДО ЗАХИСТУ', 8: 'МАГІСТЕРСЬКА КВАЛІФІКАЦІЙНА РОБОТА',
               10: 'Спеціальність', 12: 'Освітня програма', 14: 'Орієнтація освітньої програми',
               16: 'Гарант освітньої програми', 19: 'Керівник магістерської кваліфікаційної роботи',
               22: 'Виконав'}
SMALL = 9                                 # дрібні підрядкові підписи форми, пт
DATE_BLANK = '«___» ___________ 20__ р.'


def build_title(doc, lines):
    """Титульний лист за формою кафедри (методичка, Додаток В). Слова форми й дані — з full.md,
    по одному надрукованому фрагменту на рядок (assemble.py, TITLE_PAGE); тут лише розкладка та
    службові позначки форми — місце підпису, «(підпис)», «(ПІБ)», дата."""
    T = [L(n) for n in lines]
    if len(T) != 25 or any(T[i] != v for i, v in TITLE_SHAPE.items()) or not T[9].startswith('на тему ') \
            or not T[24].startswith('КИЇВ'):
        raise SystemExit('ТИТУЛКА: рядки full.md до «РЕФЕРАТ» не відповідають формі (assemble.py, TITLE_PAGE)')
    b, c, r, u = {'bold': True}, {'align': C}, {'align': R}, {'underline': True}
    para(doc, T[0], align=C, indent=0, bold=True, spacing=1)
    para(doc, T[1], align=C, indent=0, bold=True, spacing=1, before=12, after=24)
    t = grid(doc, (80, 90))
    left = [(T[2], {**b, **c, 'after': 6}), (T[3], b), ('__________  ' + T[4], {**b, 'before': 6}),
            ('(підпис)                    (ПІБ)', {'size': SMALL}), (DATE_BLANK, {'before': 4})]
    right = [(T[5], {**b, **c, 'after': 6}), (T[6], {**b, **c}), ('__________  ' + T[7], {**b, **c, 'before': 6}),
             ('(підпис)                    (ПІБ)', {**c, 'size': SMALL}), (DATE_BLANK, {**c, 'before': 4})]
    cell_lines(t.rows[0].cells[0], left); cell_lines(t.rows[0].cells[1], right)
    para(doc, T[8], align=C, indent=0, bold=True, spacing=1, before=24, after=12)
    para(doc, T[9], align=C, indent=0, bold=True, spacing=1, after=18)
    t = grid(doc, (62, 108), rows=3)
    for row, (lab, val, note) in zip(t.rows, ((T[10], T[11], '(код і найменування)'), (T[12], T[13], '(назва)'),
                                              (T[14], T[15], '(освітньо-професійна або освітньо-наукова)'))):
        cell_lines(row.cells[0], [(lab, {})])
        cell_lines(row.cells[1], [(val, {**c, **u}), (note, {**c, 'size': SMALL, 'after': 6})])
    para(doc, '', indent=0, spacing=1, after=6)
    t = grid(doc, (72, 40, 58), rows=3)
    for row, (lab, sub_, name) in zip(t.rows, ((T[16], T[17], T[18]), (T[19], T[20], T[21]), (T[22], None, T[23]))):
        cell_lines(row.cells[0], [(lab, b)] + ([(sub_, {})] if sub_ else []))
        cell_lines(row.cells[1], [('', {'rule': True}), ('(підпис)', {**c, 'size': SMALL, 'after': 14})])
        cell_lines(row.cells[2], [(name, {**b, **c, **u})])
    para(doc, T[24], align=C, indent=0, bold=True, spacing=1, before=72)
    left = [ph for _, ph, _ in PRIVATE_FIELDS if ph not in PRIVATE]
    if left:
        REPORT['placeholders'].append('Титульний лист і реферат: ' + ', '.join(left)
                                      + ' (заповнити в MagisterDocs/private/titulka.json)')


# ---------------------------------------------------------------- приватні частини
# Картка й завдання (методичка, с. 10: картка — на обкладинці, завдання — друга сторінка) —
# документи з підписами й даними наказу; їхній текст, разом зі словами форми, живе поза
# репозиторієм: MagisterDocs/private/kartka.json і zavdannia.json (власник, 2026-10-13). Тут — лише
# те, як такий файл розкласти на сторінці. Кожен запис списку «рядки» — один із:
#   {"текст": "…", "вирівнювання": "ліворуч|центр|праворуч|ширина", "жирний": true, "кегль": 14,
#    "перед_пт": 0, "після_пт": 0, "розрядка": true}
#   {"колонки": [{"текст": "…", "ширина_мм": 60, …ті самі ключі…, "лінія": true}], "після_пт": 0}
#   {"рамка": [ …записи «текст»… ]}                 — рамка довкола кількох рядків (картка)
#   {"порожніх_рядків": 2}
# «{ключ}» у тексті — значення з titulka.json; порожнє значення друкується прочерком для
# заповнення від руки.
ALIGN = {'ліворуч': LEFT, 'центр': C, 'праворуч': R, 'ширина': J}
BLANK = '_' * 24


def _fill(text):
    def one(m):
        if m.group(1) not in DATA:
            raise SystemExit(f'приватна частина: «{{{m.group(1)}}}» — такого ключа в titulka.json немає')
        return DATA[m.group(1)] or BLANK
    return re.sub(r'\{([^{}]+)\}', one, text)


def _opts(e):
    unknown = set(e) - {'текст', 'вирівнювання', 'жирний', 'кегль', 'перед_пт', 'після_пт', 'розрядка',
                        'ширина_мм', 'лінія', 'підкреслений'}
    if unknown: raise SystemExit(f'приватна частина: невідомі ключі {sorted(unknown)} у записі {e}')
    return {'align': ALIGN[e.get('вирівнювання', 'ліворуч')], 'bold': bool(e.get('жирний')),
            'size': e.get('кегль'), 'before': e.get('перед_пт', 0), 'after': e.get('після_пт', 0),
            'rule': bool(e.get('лінія')), 'underline': bool(e.get('підкреслений'))}


def _form_line(doc, e):
    o = _opts(e)
    p = para(doc, '', align=o['align'], indent=0, spacing=1, before=o['before'], after=o['after'])
    r = p.add_run(_fill(e.get('текст', ''))); r.bold = o['bold']; r.underline = o['underline']
    if o['size']: r.font.size = Pt(o['size'])
    if e.get('розрядка'):
        sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), '60'); r._r.get_or_add_rPr().append(sp)


def emit_private_part(doc, path, name):
    """Сторінка з приватного файлу. Повертає True, якщо файл є і сторінку складено."""
    if not path or not os.path.exists(path):
        REPORT['placeholders'].append(f'{name}: файла {os.path.basename(path) if path else "—"} немає — '
                                      f'сторінку не вставлено')
        return False
    spec = json.load(open(path, encoding='utf-8'))
    for e in spec['рядки']:
        if 'колонки' in e:
            cols = e['колонки']
            t = grid(doc, [c_['ширина_мм'] for c_ in cols])
            for c_, cell in zip(cols, t.rows[0].cells):
                cell_lines(cell, [(_fill(c_.get('текст', '')), _opts(c_))])
            if e.get('після_пт'): para(doc, '', indent=0, spacing=1, size=1, after=e['після_пт'])
        elif 'рамка' in e:
            # рамка — не таблиця даних: без стилю «Table Grid», за яким verify.py знаходить таблиці full.md
            t = grid(doc, (TW,))
            t._tbl.tblPr.append(parse_xml('<w:tblBorders %s>' % nsdecls('w') + ''.join(
                '<w:%s w:val="single" w:sz="6" w:space="0" w:color="000000"/>' % side
                for side in ('top', 'left', 'bottom', 'right')) + '</w:tblBorders>'))
            cell_lines(t.rows[0].cells[0], [(_fill(x.get('текст', '')), _opts(x)) for x in e['рамка']])
        elif 'порожніх_рядків' in e:
            for _ in range(e['порожніх_рядків']): para(doc, '', indent=0, spacing=1)
        else:
            _form_line(doc, e)
    REPORT['notes'].append(f'{name}: сторінку складено з {os.path.basename(path)}')
    return True


def _front_lines():
    """Номери рядків: титулка, український і англійський реферат; номер рядка «ЗМІСТ»."""
    ua, en, toc = anchor('**РЕФЕРАТ**'), anchor('**ABSTRACT**'), anchor('**ЗМІСТ**')
    title = [n for n in range(1, ua) if LINES[n - 1].strip()]
    return title, (ua, en), (en, toc), toc


def build_front(doc, parts=None):
    """Те, що стоїть перед вступом, у порядку методички (с. 10): картка (у файлі для перевірки —
    перша сторінка), титульний лист — сторінка 1, завдання — сторінка 2, зміст — сторінка 3, і
    тільки він нумерується. Реферат сюди не входить: його вкладають у конверт (build_referat).
    Повертає номер рядка, з якого продовжує build_body."""
    title, _, _, toc = _front_lines()
    card = emit_private_part(doc, parts and os.path.join(parts, 'kartka.json'), 'картка')
    if card: page_break(doc)
    build_title(doc, title)
    page_break(doc)
    if not emit_private_part(doc, parts and os.path.join(parts, 'zavdannia.json'), 'завдання'):
        boxed(doc, '[МІСЦЕ ДЛЯ ЗАВДАННЯ: друга сторінка роботи — завдання до виконання магістерської '
                   'кваліфікаційної роботи з підписами; MagisterDocs/private/zavdannia.json]')
    if card:                               # картка стоїть перед сторінкою 1: відлік із нуля
        doc.sections[0]._sectPr.append(parse_xml('<w:pgNumType %s w:start="0"/>' % nsdecls('w')))
    numbered_section(doc, 3)
    p = para(doc, L(toc).replace('**', ''), align=C, indent=0, bold=True, after=18, keep=True)
    toc_field(doc)
    return toc + 2


def build_referat(out, pages=None):
    """Реферат — окремим файлом (методичка, с. 11: вкладається в конверт разом із рецензією).
    Український і іноземною мовою, кожен з нової сторінки; текст — ті самі рядки full.md між
    титулкою і змістом."""
    _, ua, en, _ = _front_lines()
    doc = new_doc()
    for k, (head, end) in enumerate((ua, en)):
        p = para(doc, L(head).replace('**', ''), align=C, indent=0, bold=True, after=18, keep=True)
        p.paragraph_format.page_break_before = bool(k)
        for n in range(head + 1, end):
            if L(n).strip(): para(doc, L(n))
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    doc.save(out)


def build_body(doc, lo, hi):
    """Універсальний обробник: заголовки, код, формули, таблиці з підписами,
    лістинги, плейсхолдери, списки, проза. Текст — дослівно з full.md."""
    n, in_refs, last_heading = lo, False, ''
    runin = None            # заголовок пункту, що чекає на свій текст
    table_title = None      # номер таблиці, після якого наступний абзац — її назва
    in_legend = False       # ПЕРЕЛІК УМОВНИХ ПОЗНАЧЕНЬ: рядки «X – Y» збираються у дві колонки
    legend = []
    after_formula = False   # після формули вже стоїть вільний рядок
    while n <= hi:
        s = L(n)
        if not s.strip(): n += 1; continue    # поза огорожею порожній рядок — межа абзаців markdown, не код
        pending, runin = runin, None          # заголовок пункту йде в підбір лише до звичайного абзацу
        was_formula, after_formula = after_formula, False

        if legend and not (in_legend and LEGEND_ITEM.match(s)):
            emit_legend(doc, legend); legend = []     # група скінчилася: далі заголовок або проза

        m = RE_H.match(s)
        if m:
            lvl, txt = len(m.group(1)), m.group(2)
            in_refs = 'СПИСОК ВИКОРИСТАНИХ ДЖЕРЕЛ' in txt
            in_legend = lvl == 1 and txt.startswith('ПЕРЕЛІК УМОВНИХ ПОЗНАЧЕНЬ')
            ap = RE_APPENDIX.match(txt) if lvl == 2 else None
            if lvl == 1:
                struct_head(doc, txt)
            elif ap:
                # кожен додаток — з нової сторінки; перший, що йде одразу за «ДОДАТКИ», лишається
                # з цим заголовком на його сторінці
                appendix_head(doc, ap.group(1), ap.group(2), brk=not last_heading.startswith('ДОДАТКИ'))
            elif lvl == 2:
                p = doc.add_paragraph(style='Heading 2'); add_runs(p, txt)
            else:
                runin = runin_heading(doc, txt)
            last_heading = txt
            _emit_inserts(doc, n)
            n += 1; continue
        last_heading = ''

        if in_legend and LEGEND_ITEM.match(s):
            legend.append(s); n += 1; continue

        if any(a <= n <= b for a, b in CODE_RANGES):
            if not s.lstrip().startswith('```'):
                para(doc, s, indent=0, spacing=1, mono='Courier New', size=9.5, align=LEFT)
            n += 1; continue

        if s.lstrip().startswith('```'):
            n = emit_fenced(doc, n); continue

        if RE_FORMULA.match(s) and RE_FORMULA.match(s).group(2) in FORMULA_TEX:
            # методичка, с. 19: над і під формулою — не менше одного вільного рядка; між двома
            # формулами поспіль вільний рядок один, а не два
            if not was_formula: para(doc, '', indent=0, keep=True)
            emit_formula(doc, n)
            para(doc, '', indent=0)
            after_formula = True
            n += 1; continue

        if s.startswith('|'):
            n = emit_table(doc, n); continue

        if RE_TAB_CAP.match(s):
            # методичка, с. 20–21: «Таблиця N» — праворуч, назва — під ним симетрично до тексту
            p = para(doc, s, indent=0, align=R, spacing=1, before=6, keep=True, literal=True)
            table_title = CUR_TABLE[0] = RE_TAB_CAP.match(s).group(1)
            if table_title in NEW_PAGE: p.paragraph_format.page_break_before = True
            n += 1; continue

        if table_title is not None:
            p = para(doc, '', indent=0, align=C, spacing=1, after=6, keep=True)
            r = p.add_run(s); r.bold = True
            REPORT['tables'].append(f'ряд {n}: підпис «Таблиця {table_title}» + назва «{s[:50]}»')
            table_title = None
            n += 1; continue

        if RE_LST_CAP.match(s):
            para(doc, s, indent=0, spacing=1, before=6, after=6, keep=True, literal=True)
            REPORT['listings'].append(f'ряд {n}: {s[:60]}')
            k = n + 1
            while k <= hi and not L(k).strip(): k += 1
            n = emit_fenced(doc, k) if L(k).lstrip().startswith('```') else emit_listing(doc, n + 1)
            continue

        m = RE_FIG_PH.match(s)
        if m:
            if m.group(1) in FIGURES:
                emit_figure(doc, m.group(1))
            else:
                boxed(doc, s)
                REPORT['placeholders'].append(f'ряд {n}: рисунок {m.group(1)}')
            n += 1; continue

        if RE_FILL_PH.search(s):
            boxed(doc, s)
            REPORT['placeholders'].append(f'ряд {n}: {s[:60]}…')
            n += 1; continue

        if RE_FIG_CAP.match(s):
            para(doc, s, indent=0, align=C, spacing=1, after=12, literal=True)
            REPORT['figures'].append(f'ряд {n}: {s[:80]}')
            emit_pending_sheets(doc, RE_FIG_CAP.match(s).group(1))
            n += 1; continue

        if RE_NUM_ITEM.match(s) or RE_BULLET.match(s):
            para(doc, s, indent=1.25, literal=in_refs); n += 1; continue

        if pending is not None:              # текст пункту — у підбір до його заголовка
            pending.paragraph_format.keep_with_next = False
            add_runs(pending, ' ' + s)
        else:
            para(doc, s)
        _emit_inserts(doc, n)
        n += 1
    if legend: emit_legend(doc, legend)
    return doc


LEGEND_ITEM = re.compile(r'^(\S.*?) – (\S.*)$')


def emit_legend(doc, items):
    """Перелік умовних позначень двома колонками (методичка, с. 12): зліва скорочення, справа —
    розшифровка. Таблиця без рамок; текст кожного рядка — дослівно з full.md, тире лишається
    на початку правої колонки."""
    pairs = [LEGEND_ITEM.match(x).groups() for x in items]
    left = min(60, max(len(a) for a, _ in pairs) * 2.6 + 4)          # мм: найдовше скорочення в один рядок
    t = doc.add_table(rows=len(pairs), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    t._tbl.tblPr.append(parse_xml('<w:tblLayout %s w:type="fixed"/>' % nsdecls('w')))
    for j, w in enumerate((left, TW - left)):
        t.columns[j].width = Mm(w)
        for row in t.rows: row.cells[j].width = Mm(w)
    for row, (a, b) in zip(t.rows, pairs):
        for c, txt in zip(row.cells, (a, '– ' + b)):
            c.text = ''
            p = c.paragraphs[0]; f = p.paragraph_format
            f.first_line_indent = Cm(0); p.alignment = LEFT
            add_runs(p, txt)
    REPORT['notes'].append(f'перелік умовних позначень: {len(pairs)} рядків у дві колонки, ліва {left:.0f} мм')


# ================================================================ рисунки
# Рисунки рендерить render_figures.py (крок run_all.sh перед складанням) у теку з manifest.json.
# Плейсхолдер «[МІСЦЕ ДЛЯ РИСУНКА N]» заміняється зображенням; рисунок без рендера лишається
# плейсхолдером у рамці. Рисунок на кількох аркушах: перший аркуш — над підписом із full.md
# (дослівно), кожен наступний — після нього з підписом «Рис. N, аркуш k».

FIGURES, PENDING, LANDSCAPE_OPEN = {}, {}, set()
FIG_DIR = None
AREA_MM = {'portrait': (TW, 240),         # область рисунка: книжкова сторінка без підпису, мм
           'landscape': (257, 150)}       # альбомна: 297 − 20 − 20; 210 − 30 − 10 − підпис
PX_MM = 25.4 / 96                         # CSS-піксель рендера в мм за натурального розміру
RE_SHEET_LABEL = 'Рис. {num}, аркуш {k}'


def load_figures(d):
    global FIG_DIR
    FIGURES.clear(); PENDING.clear(); LANDSCAPE_OPEN.clear(); FIG_DIR = d
    if d:
        FIGURES.update(json.load(open(os.path.join(d, 'manifest.json'), encoding='utf-8')))


def _figure_scale(fig):
    """Один масштаб на всі аркуші рисунка (мм на CSS-піксель): найменший, з яким кожен аркуш
    вміщується в область сторінки, і не більший за натуральний — кегль підписів однаковий на
    всіх аркушах і дорівнює найменшому з виміряних render_figures.py."""
    W, H = AREA_MM[fig.get('orientation', 'portrait')]
    return min(min(W / s['css_px'][0], H / s['css_px'][1], PX_MM) for s in fig['sheets'])


def orientation_section(doc, landscape):
    """Розрив розділу з нової сторінки. Альбомна сторінка читається поворотом за годинниковою
    стрілкою, тож корінець (поле 30 мм) стає її верхнім полем: 30 / 10 / 20 / 20 мм, а номер
    сторінки — не вгорі біля корінця, а внизу праворуч (це правий верхній кут аркуша, як на
    книжкових). Нумерація продовжується; «особлива перша сторінка» (титулка) вимкнена.

    python-docx ставить розрив у новий порожній абзац (стиль Normal, 1,5 інтервалу, ≈ 8,5 мм):
    після підпису альбомного рисунка він міг вивести зайву порожню сторінку в Word. Розрив
    переноситься в попередній абзац; якщо перед ним таблиця — абзац лишається, але 1 пт."""
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.different_first_page_header_footer = False
    for pg in s._sectPr.findall(qn('w:pgNumType')):            # новий розділ успадкував «нумерувати з 3»:
        s._sectPr.remove(pg)                                   # нумерація має продовжуватися, а не починатися знову
    brk = doc.element.body.findall(qn('w:p'))[-1]              # новий абзац із розривом
    prev = brk.getprevious()
    if prev is not None and prev.tag == qn('w:p') and prev.find(qn('w:pPr') + '/' + qn('w:sectPr')) is None:
        prev_ppr = prev.find(qn('w:pPr'))
        if prev_ppr is None:
            prev_ppr = OxmlElement('w:pPr'); prev.insert(0, prev_ppr)
        prev_ppr.append(brk.find(qn('w:pPr') + '/' + qn('w:sectPr')))
        brk.getparent().remove(brk)
    else:
        ppr = brk.find(qn('w:pPr'))
        ppr.insert(0, parse_xml('<w:spacing %s w:before="0" w:after="0" w:line="20" w:lineRule="exact"/>' % nsdecls('w')))
    for part in (s.header, s.footer):
        part.is_linked_to_previous = False
    page_number(s.footer.paragraphs[0] if landscape else s.header.paragraphs[0])
    if landscape:
        s.orientation, s.page_width, s.page_height = WD_ORIENT.LANDSCAPE, Mm(297), Mm(210)
        s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Mm(30), Mm(10), Mm(20), Mm(20)
    else:
        s.orientation, s.page_width, s.page_height = WD_ORIENT.PORTRAIT, Mm(210), Mm(297)
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Mm(30), Mm(10), Mm(20), Mm(20)


def image_para(doc, path, width_mm):
    """Абзац з одним зображенням: по центру, одинарний інтервал (множник 1,5 додав би половину
    висоти зображення над ним), тримається з наступним — із підписом."""
    p = doc.add_paragraph(); f = p.paragraph_format
    f.first_line_indent = Cm(0); f.alignment = C; f.line_spacing_rule = WD_LINE_SPACING.SINGLE
    f.space_before = Pt(6); f.space_after = Pt(6); f.keep_with_next = True
    p.add_run().add_picture(path, width=Mm(width_mm))
    return p


def emit_image_grid(doc, num, fig):
    """Знімки (рисунок 4.1): таблиця без рамок, до чотирьох у ряд, у порядку назв файлів."""
    from docx.image.image import Image as DocxImage
    files = [os.path.join(FIG_DIR, s['file']) for s in fig['sheets']]
    per_row = min(4, len(files)); rows = -(-len(files) // per_row)
    gap = 4; cw = (TW - gap * (per_row - 1)) / per_row; max_h = 200 / rows
    t = doc.add_table(rows=rows, cols=per_row)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, f in enumerate(files):
        im = DocxImage.from_file(f)
        w = min(cw, max_h * im.px_width / im.px_height)
        c = t.rows[i // per_row].cells[i % per_row]; p = c.paragraphs[0]; pf = p.paragraph_format
        pf.first_line_indent = Cm(0); pf.alignment = C; pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.keep_with_next = True
        p.add_run().add_picture(f, width=Mm(w))
    REPORT['figures'].append(f'рисунок {num}: знімків {len(files)}, у ряд до {per_row}')


def emit_figure(doc, num):
    fig = FIGURES[num]
    if fig['kind'] == 'images':
        emit_image_grid(doc, num, fig); return
    k = _figure_scale(fig)
    if fig.get('orientation') == 'landscape':
        orientation_section(doc, True); LANDSCAPE_OPEN.add(num)
    first, rest = fig['sheets'][0], fig['sheets'][1:]
    image_para(doc, os.path.join(FIG_DIR, first['file']), first['css_px'][0] * k)
    PENDING[num] = [(s, k) for s in rest]
    pt = min(s['min_pt'] for s in fig['sheets'])
    REPORT['figures'].append(f'рисунок {num}: {len(fig["sheets"])} аркуш(і), найдрібніший підпис {pt:.1f} пт'
                             + (', альбомна сторінка' if fig.get('orientation') == 'landscape' else ''))


def emit_pending_sheets(doc, num):
    for k_sheet, (s, k) in enumerate(PENDING.pop(num, []), 2):
        image_para(doc, os.path.join(FIG_DIR, s['file']), s['css_px'][0] * k)
        para(doc, RE_SHEET_LABEL.format(num=num, k=k_sheet), indent=0, align=C, spacing=1, after=12,
             literal=True)
    if num in LANDSCAPE_OPEN:
        orientation_section(doc, False); LANDSCAPE_OPEN.discard(num)


def build(src, out, brief=None, registry=None, report=None, figures=None, private=None, parts=None,
          referat=None, pages=None, splits=None):
    load(src, brief)
    load_private(private, pages)
    load_splits(splits)
    if registry: bind_registry(registry)
    load_figures(figures)
    doc = new_doc()
    start = build_front(doc, parts)
    build_body(doc, start, len(LINES) - 1)
    if not pages:
        REPORT['placeholders'].append('кількість сторінок: не задано (--pages) — у тексті лишився плейсхолдер')
    if referat: build_referat(referat, pages)
    if PENDING:
        raise SystemExit(f'РИСУНОК БЕЗ ПІДПИСУ: аркуші {sorted(PENDING)} не мають підпису «Рис. N …» у full.md')
    for kind, items in NORMALIZED.items():
        if items: REPORT['notes'].append(f'нормалізація «{kind}»: {len(items)} замін')
    if APOSTROPHE_KV:
        REPORT['notes'].append("рядки keep-verbatim, у яких апостроф змінився з ' на ’: "
                               + '; '.join(sorted(APOSTROPHE_KV)))
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    doc.save(out)
    if report:
        json.dump(REPORT, open(report, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    return REPORT


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--src', required=True, help='full.md')
    ap.add_argument('--out', required=True, help='вихідний .docx')
    ap.add_argument('--brief', help='formatter-brief.md (для запобіжника keep-verbatim)')
    ap.add_argument('--registry', default=os.path.join(HERE, 'registry.json'))
    ap.add_argument('--report', help='куди записати звіт складання (JSON)')
    ap.add_argument('--figures', help='тека рендерів із manifest.json (render_figures.py); без неї — плейсхолдери')
    ap.add_argument('--private', help='titulka.json — дані титулки й реферату (поза репозиторієм)')
    ap.add_argument('--parts', help='тека з приватними частинами kartka.json і zavdannia.json (MagisterDocs/private)')
    ap.add_argument('--referat', help='куди записати реферат окремим .docx')
    ap.add_argument('--pages', type=int, help='кількість сторінок роботи з PDF попереднього проходу')
    ap.add_argument('--splits', help='план поділу таблиць за сторінками з PDF попереднього проходу (layout.py tables)')
    a = ap.parse_args()
    rep = build(a.src, a.out, a.brief, a.registry, a.report, a.figures, a.private, a.parts, a.referat, a.pages,
                a.splits)
    for x in rep['notes']: print('•', x)
    for x in rep['figures']:
        if x.startswith('рисунок'): print('•', x)
    print(f'формул: {len(rep["formulas"])} | таблиць: {len([t for t in rep["tables"] if "підпис" not in t])} '
          f'| плейсхолдерів: {len(rep["placeholders"])}')
    print('OK ->', a.out)
